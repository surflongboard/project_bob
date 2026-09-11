"""
Sell-Down Priority: joins the FW27 range assortment plan's planned exit
season to current warehouse available stock, so stock sitting on a
franchise that's about to leave the collection surfaces BEFORE it retires,
not after.

This is a different signal from analyze_inventory.py's Season-Recency view
(REG-INV-008/009), and deliberately complements rather than replaces it:
  - analyze_inventory.py asks "how OLD is this stock?" (backward-looking,
    from the stock file's own Season Article code on each unit).
  - This script asks "how much RUNWAY does this stock's franchise have
    left in the collection?" (forward-looking, from the assortment plan's
    own planned exit season for that style).
A franchise can be low on the first (nearly-new stock) and high-risk on
the second (already planned to exit next season) — that's exactly the
case this script is built to catch early enough to act on.

Source: data/inventory_analysis/inputs/assortment/Assortment Attribution
Review_11092026.xlsx, FW27 tab only (per the business's own framing: FW27
is "the workable season" — see ASSUMPTIONS_REGISTER.md REG-INV-010). Two
columns carry the signal, both business-identified by column letter:
  - Column AI, "LSO / Exit Season" — the planned Last Season Offered.
  - Column AN, "NEW MAPPING: Activity" — a merchandising Activity grouping
    (Day Hiking, Trekking, Active Essentials, ...), carried through for
    context/filtering, not used in the ranking itself.

Matched to franchise (Base+Gender) via ss26_lib.franchise_key() on the
FW27 tab's own `Style` text — the tab's `Franchise/Family` column is
almost entirely blank (REG-INV-011), so this is the only usable join key,
same fallback already used for the stock file itself.

CAVEATS (see ASSUMPTIONS_REGISTER.md for the full reasoning):
  - Only 392 of ~1,860 published franchises appear on the FW27 tab at
    all — this plan is scoped to FW27's own active/carry-over/new range,
    not the full historical franchise universe. Stock on a franchise with
    no FW27-tab row is reported separately (Sheet 4), not silently
    dropped and not assumed safe.
  - Sheet 4 is NOT safe to blanket-label "retired before FW27" (REG-INV-012):
    it's cross-checked against the SS27 tab (the season immediately before
    FW27), and a real share of it was still Active/Carry-Over there. That
    subset needs a merch check, not an assumption — see the sheet's own
    flag column and REG-INV-012 for two confirmed examples where the
    likely explanation is a name change (e.g. "Rosson Softshell Hood" ->
    "Rosson Mid II Hood"), not a real exit, which Base+Gender matching
    can't catch (a known limitation, documented in config.py itself).
  - Sheet 5 (REG-INV-013) narrows that subset further with a heuristic
    rename-candidate finder — a ranked suggestion list to check first, not
    a resolved mapping. It never feeds back into Sheets 2-4's numbers.
  - Sheets 2/3/6 (REG-INV-014) add an inventory value estimate (Retail/
    Wholesale/Est. Cost). Landed Cost is broken in the source (#ERROR! on
    every FW27 row) and is reconstructed as WP x (1-GM0%) instead --
    validated against SS27's own non-broken figures, but Target RRP/WP
    themselves are planned prices, not realized/audited ones. Sheet 4 has
    no price source at all and is excluded from every value total (shown
    as units-only, not zero).
  - "REVIEW" (Exit Season undecided) is NOT force-ranked into the sell-
    down priority list — REG-INV-010 documents this as a genuine Open
    item per this repo's register discipline. Those franchises get their
    own sheet (Sheet 3) so the stock at stake is visible without
    pretending the timing is known.
  - "FW27 or SS28" (an explicitly ambiguous exit season in the source
    data) is kept as its own tier rather than collapsed into either
    season — see REG-INV-010.

Usage:
    python3 data/inventory_analysis/scripts/build_exit_plan_priority.py \
        [--assortment PATH] [--stock PATH]
"""
from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font

REPO_ROOT = Path(__file__).resolve().parents[3]
sys.path.insert(0, str(REPO_ROOT / "data" / "ss26_portfolio_tiering" / "scripts"))
import ss26_lib as lib  # noqa: E402

WORKSTREAM_DIR = REPO_ROOT / "data" / "inventory_analysis"
DEFAULT_ASSORTMENT = WORKSTREAM_DIR / "inputs" / "assortment" / "Assortment Attribution Review_11092026.xlsx"
DEFAULT_STOCK = REPO_ROOT / "data" / "ss26_portfolio_tiering" / "inputs" / "stock" / "Available_stock_260825.xlsx"
ASSORTMENT_SHEET = "FW27"
PRIOR_SEASON_SHEET = "SS27"  # cross-check for Sheet 4 -- see REG-INV-012

# --------------------------------------------------------------------------
# REG-INV-010: exit-urgency tiers, read directly off the FW27 tab's own
# "LSO / Exit Season" values (column AI) — only these five values exist in
# the current pull, so this is an explicit lookup, not season arithmetic.
# Ranked list order = TIER_RANK order; "Pending decision" is NEVER placed
# in the ranked list (see module docstring).
# --------------------------------------------------------------------------
TIER_RANK = {
    "FW27": (1, "Exiting FW27 (last season offered)"),
    "FW27 or SS28": (2, "Exiting FW27 or SS28 (ambiguous — treat as near-term)"),
    "SS28": (3, "One season left (exits after SS28)"),
    "FW28+": (4, "Long runway (continues through FW28+)"),
}
PENDING_VALUE = "REVIEW"

# --------------------------------------------------------------------------
# REG-INV-013: heuristic name-change candidate finder for the "still Active
# in SS27" subset of Sheet 4 (REG-INV-012). This does NOT resolve REG-INV-012
# -- it only narrows "needs a merch check" down to "here's a plausible
# successor to check first," ranked by confidence. Never auto-applied to
# any figure elsewhere in this workstream.
#
# Method: franchise_key() already catches a rename that's just a gender/
# version-token difference (that's not what lands in Sheet 4 at all -- those
# match already). What's left in Sheet 4 is a genuine text change, so this
# looks for shared DISTINCTIVE words between the orphaned SS27 style and
# every FW27 style of the same gender and Business Area.
#
# "Distinctive" is decided positionally, not by raw frequency: a product
# family name (Rosson, Korp, Zircon, ...) almost always LEADS the style
# name (or follows a brand/sub-brand prefix like "Haglöfs"/"L.I.M."), while
# a generic descriptor or garment noun (Jacket, GTX, Proof, Mid, Hood, ...)
# rarely does, regardless of how often it appears overall -- e.g. "Rosson"
# and "Mid" both occur ~15-50 times across FW27, but only "Rosson" is a
# family name; raw frequency alone can't tell them apart, leading position
# can. BRAND_STOP excludes the two house-wide prefixes explicitly, since
# they lead almost everything and would otherwise look "distinctive" too.
#
# Confidence: HIGH if the two names' last word (garment noun, e.g. "Hood"/
# "Pant") also matches -- same family AND same silhouette. MEDIUM if not,
# but 2+ distinctive words are shared. LOW if exactly one distinctive word
# is shared and the garment noun differs (weak family-line overlap only).
# All ties at the best confidence level are surfaced (up to 3), not just
# one -- collapsing to a single guess would hide equally-plausible options
# from the human doing the actual check.
# --------------------------------------------------------------------------
BRAND_STOP_WORDS = {"haglöfs", "l.i.m", "af", "ii", "iii", "iv", "2.0", "1", "2"}


def _words(base: str) -> list[str]:
    return [w.lower() for w in re.findall(r"[\w.]+", base, flags=re.UNICODE) if w]


def _leading_word(ws: list[str]) -> str | None:
    if not ws:
        return None
    if ws[0] in BRAND_STOP_WORDS and len(ws) > 1:
        return ws[1]
    return ws[0]


def build_distinctiveness(fw27_bases: list[str]):
    """Returns is_distinctive(word) -> bool, fit on the FW27 tab's own base names."""
    lead_count, total_count = Counter(), Counter()
    for base in fw27_bases:
        ws = _words(base)
        if not ws:
            continue
        for w in set(ws):
            total_count[w] += 1
        lead_count[_leading_word(ws)] += 1

    def is_distinctive(word: str) -> bool:
        if word in BRAND_STOP_WORDS:
            return False
        tc = total_count[word]
        if tc <= 1:
            return True
        return lead_count[word] / tc >= 0.7

    return is_distinctive


def find_rename_candidates(orphan_base, business_area, gender, fw27_by_gender_area, is_distinctive):
    """Up to 3 same-confidence FW27 base names most likely to be orphan_base's
    successor under a new name. Returns [] if nothing shares a distinctive word."""
    wa = _words(orphan_base)
    sa, last_a = set(wa), (wa[-1] if wa else None)
    candidates = []
    for fw_base in fw27_by_gender_area.get((gender, business_area), []):
        wb = _words(fw_base)
        shared = [w for w in (sa & set(wb)) if is_distinctive(w)]
        if not shared:
            continue
        last_b = wb[-1] if wb else None
        tier = "HIGH" if last_a == last_b else ("MEDIUM" if len(shared) >= 2 else "LOW")
        candidates.append((tier, len(shared), fw_base, shared))
    if not candidates:
        return []
    tier_rank = {"HIGH": 2, "MEDIUM": 1, "LOW": 0}
    best_rank = max(tier_rank[c[0]] for c in candidates)
    best = sorted((c for c in candidates if tier_rank[c[0]] == best_rank), key=lambda c: -c[1])
    return best[:3]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--assortment", default=str(DEFAULT_ASSORTMENT))
    ap.add_argument("--stock", default=str(DEFAULT_STOCK))
    args = ap.parse_args()

    # --- FW27 exit plan ---
    wb_a = openpyxl.load_workbook(args.assortment, data_only=True)
    ws_a = wb_a[ASSORTMENT_SHEET]
    rows = list(ws_a.iter_rows(values_only=True))
    hdr = rows[1]
    idx = {h: i for i, h in enumerate(hdr) if h}
    for col in ("Code", "Style", "LSO / Exit Season ", "NEW MAPPING: Activity", "Active", "Business Area",
                "Target RRP", "WP", "GM0%"):
        if col not in idx:
            sys.exit(f"Expected column '{col}' not found on {ASSORTMENT_SHEET} header: {hdr}")
    plan = {}
    fw27_by_gender_area = defaultdict(list)
    fw27_bases = []
    n_cost_reconstructed = 0
    for r in rows[2:]:
        if r[idx["Code"]] in (None, ""):
            continue
        style = str(r[idx["Style"]])
        base, gender = lib.franchise_key(style)
        rrp, wp, gm = r[idx["Target RRP"]], r[idx["WP"]], r[idx["GM0%"]]
        # REG-INV-014: LANDED COST SEK itself is #ERROR! on every FW27 row (REG-INV-011).
        # Reconstructed instead as WP x (1 - GM0%) -- algebraically what GM0% means,
        # validated against SS27's own (non-broken) LANDED COST SEK where both are
        # present: exact match to floating-point noise. None when GM0% isn't numeric
        # (5 of 392 rows are #DIV/0!) rather than a guessed cost.
        cost = wp * (1 - gm) if isinstance(wp, (int, float)) and isinstance(gm, (int, float)) else None
        if cost is not None:
            n_cost_reconstructed += 1
        plan[(base, gender)] = {
            "style": style,
            "active": r[idx["Active"]],
            "exit_season": r[idx["LSO / Exit Season "]],
            "activity": r[idx["NEW MAPPING: Activity"]],
            "business_area": r[idx["Business Area"]],
            "rrp": rrp if isinstance(rrp, (int, float)) else None,
            "wp": wp if isinstance(wp, (int, float)) else None,
            "cost": cost,
        }
        fw27_by_gender_area[(gender, r[idx["Business Area"]])].append(base)
        fw27_bases.append(base)
    print(f"FW27 exit-plan rows: {len(plan)}  (cost reconstructed for {n_cost_reconstructed})")
    is_distinctive = build_distinctiveness(fw27_bases)

    # --- SS27 tab, for the Sheet-4 "actually retired?" cross-check (REG-INV-012) ---
    ws_prior = wb_a[PRIOR_SEASON_SHEET]
    prows = list(ws_prior.iter_rows(values_only=True))
    phdr = prows[1]
    pidx = {h: i for i, h in enumerate(phdr) if h}
    prior_season = {}
    for r in prows[2:]:
        if r[pidx["Code"]] in (None, ""):
            continue
        key = lib.franchise_key(str(r[pidx["Style"]]))
        prior_season[key] = {
            "active": r[pidx["Active"]],
            "status": r[pidx.get("STATUS", -1)],
            "business_area": r[pidx.get("Business Area", -1)],
        }
    print(f"{PRIOR_SEASON_SHEET} rows (for cross-check): {len(prior_season)}")

    # --- current stock, same matching as analyze_inventory.py / update_stock.py ---
    wb_s = openpyxl.load_workbook(args.stock, data_only=True, read_only=True)
    candidates = [n for n in wb_s.sheetnames if n.lower().startswith("available stock")]
    ws_s = wb_s[candidates[0]] if candidates else wb_s[wb_s.sheetnames[0]]
    srows = list(ws_s.iter_rows(values_only=True))
    shdr = srows[0]
    sidx = {h: i for i, h in enumerate(shdr)}
    stock_units = defaultdict(float)
    for r in srows[2:]:
        article = r[sidx["Article"]]
        units = r[sidx["Available stock"]] or 0
        if not article:
            continue
        stock_units[lib.franchise_key(str(article).strip())] += units
    print(f"stock franchises with units: {len(stock_units)}")

    # --- join ---
    ranked = []       # franchises with stock, matched to a ranked exit tier
    pending = []       # franchises with stock, exit season = REVIEW
    unplanned = []     # franchises with stock, no FW27-tab row at all
    for key, units in stock_units.items():
        if units <= 0:
            continue
        info = plan.get(key)
        if info is None:
            unplanned.append((key, units))
            continue
        exit_season = info["exit_season"]
        if exit_season == PENDING_VALUE:
            pending.append((key, units, info))
        elif exit_season in TIER_RANK:
            tier_num, tier_label = TIER_RANK[exit_season]
            ranked.append((key, units, info, tier_num, tier_label))
        else:
            # An exit-season value outside the five known ones -- don't guess
            # its urgency, route it alongside "pending" so it's visible.
            pending.append((key, units, {**info, "exit_season": f"UNRECOGNIZED: {exit_season!r}"}))

    ranked.sort(key=lambda x: (x[3], -x[1]))
    pending.sort(key=lambda x: -x[1])
    unplanned.sort(key=lambda x: -x[1])

    print(f"matched to a ranked exit tier: {len(ranked)}")
    print(f"pending merch decision (REVIEW): {len(pending)}")
    print(f"no FW27-tab row at all: {len(unplanned)}")
    still_active_ss27 = [(k, u) for k, u in unplanned if prior_season.get(k, {}).get("active") is True]
    print(f"  of which still Active in {PRIOR_SEASON_SHEET}: {len(still_active_ss27)} "
          f"({sum(u for k, u in still_active_ss27):,.0f} units) -- needs a merch check, not an assumption")

    # --- REG-INV-013: rename candidates for the still-Active-in-SS27 subset ---
    rename_candidates = []
    for key, units in still_active_ss27:
        base, gender = key
        business_area = prior_season[key]["business_area"]
        cands = find_rename_candidates(base, business_area, gender, fw27_by_gender_area, is_distinctive)
        rename_candidates.append((key, units, cands))
    n_with_candidate = sum(1 for _, _, c in rename_candidates if c)
    print(f"  of those, found a plausible FW27 rename candidate for: {n_with_candidate} "
          f"of {len(still_active_ss27)} -- still needs human confirmation, not auto-applied")

    # --- REG-INV-014: value rollups (Sell-Down Priority + Pending only -- Sheet 4 has no price source) ---
    valued_rows = [(key, units, info) for key, units, info, *_ in ranked] + [(key, units, info) for key, units, info in pending]
    total_retail_value = sum(u * info["rrp"] for _, u, info in valued_rows if info["rrp"] is not None)
    total_wholesale_value = sum(u * info["wp"] for _, u, info in valued_rows if info["wp"] is not None)
    total_cost_value = sum(u * info["cost"] for _, u, info in valued_rows if info["cost"] is not None)
    n_valued_units = sum(u for _, u, info in valued_rows if info["cost"] is not None)

    total_stock = sum(stock_units.values())
    print(f"total stock units: {total_stock:,.0f}  "
          f"ranked: {sum(x[1] for x in ranked):,.0f}  "
          f"pending: {sum(x[1] for x in pending):,.0f}  "
          f"unplanned: {sum(x[1] for x in unplanned):,.0f}")

    # ---------------------------------------------------------------------
    # Build the workbook
    # ---------------------------------------------------------------------
    s = lib.styles()
    HDR_FONT, HDR_FILL, BODY_FONT, GRAY_FONT = s["header_font"], s["header_fill"], s["body_font"], s["gray_font"]
    TITLE_FONT = Font(name="Arial", bold=True, color="FF1F3864")
    WRAP = Alignment(vertical="top", wrap_text=True)
    NUM_FMT = s["num_fmt"]

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws1 = wb.create_sheet("1. Overview")
    ws1["A1"] = "Sell-Down Priority — FW27 Exit Plan x Current Stock"
    ws1["A1"].font = TITLE_FONT
    ws1.merge_cells("A1:F1")
    ws1["A3"] = (
        "Joins the FW27 range assortment plan's own planned exit season (\"LSO / Exit Season\", "
        "column AI) to current warehouse stock, so a franchise already planned to leave the "
        "collection surfaces here BEFORE it retires -- not discovered afterward as dead stock. "
        "Matched to franchise (Base+Gender) via ss26_lib.franchise_key() on the assortment plan's "
        "own Style text (its Franchise/Family column is not usable -- see REG-INV-011)."
    )
    ws1["A3"].font = BODY_FONT
    ws1["A3"].alignment = WRAP
    ws1.merge_cells("A3:F3")
    ws1.row_dimensions[3].height = 75
    ws1["A6"] = (
        f"{total_stock:,.0f} total stock units across all franchises with any available stock. "
        f"{sum(x[1] for x in ranked):,.0f} matched a franchise with a known FW27 exit-season tier "
        f"(Sheet 2). {sum(x[1] for x in pending):,.0f} matched a franchise still at \"REVIEW\" -- "
        "exit timing undecided, NOT force-ranked (Sheet 3). "
        f"{sum(x[1] for x in unplanned):,.0f} has no FW27-tab row at all (Sheet 4) -- of which "
        f"{sum(u for k, u in still_active_ss27):,.0f} units ({len(still_active_ss27)} franchises) were "
        f"still Active in {PRIOR_SEASON_SHEET} and are NOT safe to assume retired (REG-INV-012)."
    )
    ws1["A6"].font = GRAY_FONT
    ws1["A6"].alignment = WRAP
    ws1.merge_cells("A6:F6")
    ws1.row_dimensions[6].height = 75
    ws1["A9"] = (
        f"Value estimate (REG-INV-014, Sheets 2-3 only -- {n_valued_units:,.0f} of {total_stock:,.0f} total "
        f"units, {n_valued_units/total_stock*100:.0f}%): {total_retail_value:,.0f} SEK at Target RRP, "
        f"{total_wholesale_value:,.0f} SEK at WP, {total_cost_value:,.0f} SEK estimated landed cost "
        "(reconstructed -- see Sheet 2's own note and REG-INV-014). Sheet 4 (Not on FW27 Plan) has no "
        "price source anywhere in this data and is NOT included in these totals -- not zero, just unknown. "
        "See Sheet 6 for the breakdown by priority tier and by Activity."
    )
    ws1["A9"].font = GRAY_FONT
    ws1["A9"].alignment = WRAP
    ws1.merge_cells("A9:F9")
    ws1.row_dimensions[9].height = 75
    for col, w in zip("ABCDEF", (20,) * 6):
        ws1.column_dimensions[col].width = w

    def write_table(ws, headers, col_widths, title, note=None):
        ws["A1"] = title
        ws["A1"].font = TITLE_FONT
        ws.merge_cells(f"A1:{openpyxl.utils.get_column_letter(len(headers))}1")
        start = 3
        if note:
            ws["A3"] = note
            ws["A3"].font = GRAY_FONT
            ws["A3"].alignment = WRAP
            ws.merge_cells(f"A3:{openpyxl.utils.get_column_letter(len(headers))}3")
            ws.row_dimensions[3].height = 45
            start = 5
        for c, h in enumerate(headers, start=1):
            cell = ws.cell(row=start, column=c, value=h)
            cell.font, cell.fill = HDR_FONT, HDR_FILL
        for c, w in enumerate(col_widths, start=1):
            ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width = w
        return start + 1

    def value_cells(ws, row, start_col, units, info):
        """Writes Retail/Wholesale/Cost value cells (units x per-unit price), blank if unpriced."""
        rrp, wp, cost = info["rrp"], info["wp"], info["cost"]
        for i, price in enumerate((rrp, wp, cost)):
            col = start_col + i
            val = round(units * price) if price is not None else None
            c = ws.cell(row=row, column=col, value=val)
            c.number_format, c.font = NUM_FMT, BODY_FONT

    VALUE_HEADERS = ["Retail Value (RRP, SEK)", "Wholesale Value (WP, SEK)", "Est. Cost Value (SEK)"]
    VALUE_WIDTHS = (20, 20, 20)
    VALUE_NOTE = (
        "Value columns per REG-INV-014: Retail/Wholesale Value = units x the FW27 plan's own "
        "Target RRP/WP for that franchise; Est. Cost Value = units x Landed Cost RECONSTRUCTED as "
        "WP x (1-GM0%) since the plan's own Landed Cost field is broken (#ERROR!, REG-INV-011). "
        "One planned price per franchise applied uniformly across all sizes/colors in stock -- not "
        "a per-SKU actual. Target RRP is a planned price, not realized/discounted revenue -- read "
        "Retail Value as a ceiling, not expected recovery."
    )

    # --- Sheet 2: ranked sell-down priority ---
    ws2 = wb.create_sheet("2. Sell-Down Priority")
    r = write_table(
        ws2,
        ["Priority Tier", "Base", "Gender", "Activity", "Exit Season", "Units in Stock", "Active?"] + VALUE_HEADERS,
        (34, 30, 10, 20, 16, 16, 10) + VALUE_WIDTHS,
        "Sell-Down Priority (most urgent exit, most stock, first)",
        "Sorted by exit urgency (Sheet 1's tier order), then by units in stock, descending. "
        "\"FW28+\" rows are long-runway, low-urgency -- included for completeness, not action. "
        + VALUE_NOTE,
    )
    for key, units, info, tier_num, tier_label in ranked:
        base, g = key
        ws2.cell(row=r, column=1, value=tier_label).font = BODY_FONT
        ws2.cell(row=r, column=2, value=base).font = BODY_FONT
        ws2.cell(row=r, column=3, value=g).font = BODY_FONT
        ws2.cell(row=r, column=4, value=info["activity"] or "").font = BODY_FONT
        ws2.cell(row=r, column=5, value=info["exit_season"]).font = BODY_FONT
        c = ws2.cell(row=r, column=6, value=round(units)); c.number_format, c.font = NUM_FMT, BODY_FONT
        ws2.cell(row=r, column=7, value=bool(info["active"])).font = BODY_FONT
        value_cells(ws2, r, 8, units, info)
        r += 1

    # --- Sheet 3: pending merch decision ---
    ws3 = wb.create_sheet("3. Pending Review")
    r = write_table(
        ws3,
        ["Base", "Gender", "Activity", "Exit Season (raw)", "Units in Stock", "Active?"] + VALUE_HEADERS,
        (34, 10, 20, 22, 16, 10) + VALUE_WIDTHS,
        "Pending Merch Decision — Exit Timing Not Yet Set",
        "These franchises carry real stock but the FW27 plan hasn't fixed an exit season yet "
        "(\"REVIEW\"). Not ranked above -- ranking these would mean guessing a business decision "
        "that hasn't been made (REG-INV-010). Sorted by units in stock, descending. " + VALUE_NOTE,
    )
    for key, units, info in pending:
        base, g = key
        ws3.cell(row=r, column=1, value=base).font = BODY_FONT
        ws3.cell(row=r, column=2, value=g).font = BODY_FONT
        ws3.cell(row=r, column=3, value=info["activity"] or "").font = BODY_FONT
        ws3.cell(row=r, column=4, value=info["exit_season"]).font = BODY_FONT
        c = ws3.cell(row=r, column=5, value=round(units)); c.number_format, c.font = NUM_FMT, BODY_FONT
        ws3.cell(row=r, column=6, value=bool(info["active"])).font = BODY_FONT
        value_cells(ws3, r, 7, units, info)
        r += 1

    # --- Sheet 4: unplanned stock (no FW27-tab row), cross-checked against SS27 ---
    still_active_prior = sum(1 for k, u in unplanned if prior_season.get(k, {}).get("active") is True)
    still_active_prior_units = sum(u for k, u in unplanned if prior_season.get(k, {}).get("active") is True)
    ws4 = wb.create_sheet("4. Not on FW27 Plan")
    r = write_table(
        ws4,
        ["Base", "Gender", "Units in Stock", f"In {PRIOR_SEASON_SHEET} Tab?", f"{PRIOR_SEASON_SHEET} Active?", f"{PRIOR_SEASON_SHEET} Status"],
        (34, 10, 16, 14, 14, 12),
        "Stock With No Row on the FW27 Tab",
        "NOT safe to blanket-label \"retired before FW27\" (REG-INV-012): cross-checked against "
        f"{PRIOR_SEASON_SHEET} (the season immediately before FW27) below. "
        f"{still_active_prior} franchises ({still_active_prior_units:,.0f} units) were still Active "
        f"there and need a real merch check -- likely explanations include a genuine drop, but also "
        "a name change too large for Base+Gender matching to catch (e.g. \"Rosson Softshell Hood\" -> "
        f"\"Rosson Mid II Hood\", confirmed both exist, neither name matches the other). Rows with "
        f"\"{PRIOR_SEASON_SHEET} Active? = False\" or blank (not in {PRIOR_SEASON_SHEET} either) are the "
        "safer default for \"likely already retired.\" Sorted by units in stock, descending.",
    )
    for key, units in unplanned:
        base, g = key
        prior = prior_season.get(key)
        ws4.cell(row=r, column=1, value=base).font = BODY_FONT
        ws4.cell(row=r, column=2, value=g).font = BODY_FONT
        c = ws4.cell(row=r, column=3, value=round(units)); c.number_format, c.font = NUM_FMT, BODY_FONT
        ws4.cell(row=r, column=4, value="Yes" if prior else "No").font = BODY_FONT
        ws4.cell(row=r, column=5, value=(prior["active"] if prior else None)).font = BODY_FONT
        ws4.cell(row=r, column=6, value=(prior["status"] if prior else None)).font = BODY_FONT
        r += 1

    # --- Sheet 5: possible FW27 renames, for the still-Active-in-SS27 subset ---
    ws5 = wb.create_sheet("5. Possible FW27 Renames")
    r = write_table(
        ws5,
        ["Base (SS27)", "Gender", "Units in Stock", "Candidate FW27 Name(s)", "Confidence", "Shared Word(s)"],
        (30, 10, 16, 46, 12, 20),
        "Heuristic Rename Candidates (REG-INV-013) — Needs Human Confirmation",
        f"For the {len(still_active_ss27)} franchises on Sheet 4 that were still Active in "
        f"{PRIOR_SEASON_SHEET}: does an FW27 style of the same gender + Business Area share a "
        "distinctive family word (e.g. \"Rosson\", \"Korp\") with this SS27 style? HIGH = same "
        "family word AND same last word (garment type). MEDIUM = 2+ shared distinctive words, "
        "different garment type. LOW = exactly one shared distinctive word, different garment "
        "type. This is a candidate list to check, not a resolved mapping -- REG-INV-012 stays "
        "Open regardless of confidence shown here. Rows with no candidate are omitted.",
    )
    for key, units, cands in sorted(rename_candidates, key=lambda x: -x[1]):
        if not cands:
            continue
        base, g = key
        names = "; ".join(c[2] for c in cands)
        shared = "; ".join(", ".join(c[3]) for c in cands)
        tier = cands[0][0]
        ws5.cell(row=r, column=1, value=base).font = BODY_FONT
        ws5.cell(row=r, column=2, value=g).font = BODY_FONT
        c = ws5.cell(row=r, column=3, value=round(units)); c.number_format, c.font = NUM_FMT, BODY_FONT
        ws5.cell(row=r, column=4, value=names).font = BODY_FONT
        ws5.cell(row=r, column=5, value=tier).font = BODY_FONT
        ws5.cell(row=r, column=6, value=shared).font = BODY_FONT
        r += 1

    # --- Sheet 6: value summary, by Priority Tier and by Activity (REG-INV-014) ---
    ws6 = wb.create_sheet("6. Inventory Value Summary")
    r = write_table(
        ws6,
        ["Priority Tier", "Units", "Retail Value (RRP, SEK)", "Wholesale Value (WP, SEK)", "Est. Cost Value (SEK)"],
        (34, 14, 20, 20, 20),
        "Inventory Value by Priority Tier",
        "Sums Sheets 2-3's value columns per tier (REG-INV-014). \"Not on FW27 Plan\" (Sheet 4) has "
        "no price source and is shown separately below with units only, not zero SEK.",
    )
    by_tier = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])  # units, retail, wholesale, cost
    for _, units, info, _, tier_label in ranked:
        row = by_tier[tier_label]
        row[0] += units
        row[1] += units * info["rrp"] if info["rrp"] is not None else 0
        row[2] += units * info["wp"] if info["wp"] is not None else 0
        row[3] += units * info["cost"] if info["cost"] is not None else 0
    pending_row = by_tier["Pending Review (exit season = REVIEW)"]
    for _, units, info in pending:
        pending_row[0] += units
        pending_row[1] += units * info["rrp"] if info["rrp"] is not None else 0
        pending_row[2] += units * info["wp"] if info["wp"] is not None else 0
        pending_row[3] += units * info["cost"] if info["cost"] is not None else 0
    tier_order_for_summary = [t[1] for t in TIER_RANK.values()] + ["Pending Review (exit season = REVIEW)"]
    for tier_label in tier_order_for_summary:
        vals = by_tier.get(tier_label)
        if not vals:
            continue
        ws6.cell(row=r, column=1, value=tier_label).font = BODY_FONT
        for c, v in enumerate(vals, start=2):
            cell = ws6.cell(row=r, column=c, value=round(v)); cell.number_format, cell.font = NUM_FMT, BODY_FONT
        r += 1
    unplanned_units = sum(x[1] for x in unplanned)
    ws6.cell(row=r, column=1, value="Not on FW27 Plan (Sheet 4 — no price source)").font = GRAY_FONT
    c = ws6.cell(row=r, column=2, value=round(unplanned_units)); c.number_format, c.font = NUM_FMT, GRAY_FONT
    ws6.cell(row=r, column=3, value="n/a").font = GRAY_FONT
    ws6.cell(row=r, column=4, value="n/a").font = GRAY_FONT
    ws6.cell(row=r, column=5, value="n/a").font = GRAY_FONT
    r += 2

    hdr_row = r
    headers2 = ["Activity", "Units", "Retail Value (RRP, SEK)", "Wholesale Value (WP, SEK)", "Est. Cost Value (SEK)"]
    ws6.cell(row=hdr_row, column=1, value="Inventory Value by Activity (merchandising category, column AN)").font = TITLE_FONT
    ws6.merge_cells(f"A{hdr_row}:E{hdr_row}")
    ws6.cell(row=hdr_row + 1, column=1, value=(
        "Same Sheets 2-3 rows, grouped by the FW27 plan's \"NEW MAPPING: Activity\" field (column AN) "
        "instead of tier. Sorted by Retail Value, descending."
    )).font = GRAY_FONT
    ws6.cell(row=hdr_row + 1, column=1).alignment = WRAP
    ws6.merge_cells(f"A{hdr_row + 1}:E{hdr_row + 1}")
    hdr_row += 3
    for c, h in enumerate(headers2, start=1):
        cell = ws6.cell(row=hdr_row, column=c, value=h)
        cell.font, cell.fill = HDR_FONT, HDR_FILL
    r = hdr_row + 1
    by_activity = defaultdict(lambda: [0.0, 0.0, 0.0, 0.0])
    for _, units, info in valued_rows:
        act = info["activity"] or "(blank)"
        row = by_activity[act]
        row[0] += units
        row[1] += units * info["rrp"] if info["rrp"] is not None else 0
        row[2] += units * info["wp"] if info["wp"] is not None else 0
        row[3] += units * info["cost"] if info["cost"] is not None else 0
    for act, vals in sorted(by_activity.items(), key=lambda kv: -kv[1][1]):
        ws6.cell(row=r, column=1, value=act).font = BODY_FONT
        for c, v in enumerate(vals, start=2):
            cell = ws6.cell(row=r, column=c, value=round(v)); cell.number_format, cell.font = NUM_FMT, BODY_FONT
        r += 1

    out_name = f"Project_Bob_Sell_Down_Priority_{date.today().strftime('%d%m%Y')}.xlsx"
    out_path = WORKSTREAM_DIR / out_name
    wb.save(out_path)
    print("saved:", out_path)


if __name__ == "__main__":
    main()
