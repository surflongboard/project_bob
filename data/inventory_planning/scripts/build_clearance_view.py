"""
Builds the clearance view for the end-of-life stock (styles NOT in the
active FW27/SS28 range, end of life at end of SS27 per IP-003): how old
each colour is, whether its style sells all year or only in one season,
what it's worth at landed cost and wholesale price, how much markdown
headroom it has before selling below cost, and how long it would take
to sell at its 2025 rate of sale -- to gauge how much of it needs
discounting to clear before FW27 starts (1-Sep-2027, IP-007).

Grain: style + colour (Article No.[:9]). Each colour carries exactly one
"Season Article" tag in the stock file -- a style's colours can be
bought in different seasons, but a single colour never is -- so age
belongs to the colour, not the style (IP-006).

Age buckets, relative to the SS27 sell-off (IP-006):
  Current (S27/F27) / Last season (F26) / 1 year (S26) /
  2 years (S25/F25) / 3+ years (F24 and older)

Style seasonality (All year / Fall/Winter / Spring/Summer), resolved in
this order, with the source kept on every row (IP-006):
  1. the code's own row in the FW27/SS28 range or Pic Centric files
     (listed there, just not active)
  2. the active successor in the same franchise (same base_name())
  3. own colours tagged to both an S and an F season -> All year (inferred)
  4. layer default: the majority Style Seasonality of that layer's styles
     in the FW27/SS28 range files -- Provisional placeholder, user-chosen
     30-Oct-2026 instead of asking Merch for the real values

Value: Available Stock units x the SKU's unit landed cost / unit WP, where
the unit figures are the quantity-weighted average across the Physical
stock and Purchase SS27 sheets (Available stock itself carries units
only). Markdown headroom = 1 - landed cost / WP: the deepest discount
off WP before selling below cost. It's a ratio of list prices, not a
sales-based margin, so the MIN_RELIABLE threshold isn't applied to it
(IP-008).

Rate of sale: 2025 Units Sold (all channels) per article name from
bob_salesdata_2025.xlsx / 52, at style level (sales data has no colour).
Weeks of cover = style Available Stock / weekly rate of sale, blanked
where the style's 2025 sales are under MIN_RELIABLE (50,000 SEK) per
CLAUDE.md -- a small denominator makes cover meaningless (IP-008).
Compared against two window lengths, since when the sell-off starts is
still Open (IP-007).

No discount % is assigned -- headroom only, per the user (30-Oct-2026).

Run: python3 build_clearance_view.py   (after build_inventory_review.py's
inputs are current -- this imports its build() directly)
Output: ../outputs/SS27_Clearance_View.xlsx
"""
from __future__ import annotations

import sys
from collections import Counter, defaultdict
from datetime import date
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

IP_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = IP_DIR.parents[1]

sys.path.insert(0, str(REPO_ROOT / "data" / "range_review" / "scripts"))
from build_franchise_list import FW27_PATH, SS28_PATH, strip_special_edition  # noqa: E402

sys.path.insert(0, str(REPO_ROOT / "data" / "ss26_portfolio_tiering" / "scripts"))
import ss26_lib as lib  # noqa: E402

sys.path.insert(0, str(Path(__file__).resolve().parent))
from build_inventory_review import STOCK_PATH, build as build_inventory  # noqa: E402

RR_INPUTS = REPO_ROOT / "data" / "range_review" / "inputs"
PIC_PATHS = [RR_INPUTS / "images" / "FW27_Pic_Centric_28_oct.xlsx",
             RR_INPUTS / "images" / "SS28_Pic_Centric_28_oct.xlsx"]
SALES_PATH = REPO_ROOT / "data" / "bob_salesdata_2025.xlsx"
OUTPUTS_DIR = IP_DIR / "outputs"

# IP-007: sell-off ends when FW27 starts; the start is still Open, so both shown.
WINDOW_END = date(2027, 9, 1)
WINDOWS = {
    "from 30-Oct-2026": date(2026, 10, 30),
    "from 1-Feb-2027 (SS27 only)": date(2027, 2, 1),
}

SEASON_ORDER = ["S22", "F22", "S23", "F23", "S24", "F24",
                "S25", "F25", "S26", "F26", "S27", "F27"]
AGE_BUCKETS = ["Current (S27/F27)", "Last season (F26)", "1 year (S26)",
               "2 years (S25/F25)", "3+ years (F24 and older)"]


def age_bucket(season: str | None) -> str:
    if season in ("S27", "F27"):
        return AGE_BUCKETS[0]
    if season == "F26":
        return AGE_BUCKETS[1]
    if season == "S26":
        return AGE_BUCKETS[2]
    if season in ("S25", "F25"):
        return AGE_BUCKETS[3]
    return AGE_BUCKETS[4]


def _norm(text) -> str:
    return " ".join(str(text or "").split())


def _layer_key(layer) -> str:
    """Stock file says "Mid layer"/"Sleepingbags", range files "Midlayer"/"Sleeping Bags"."""
    return str(layer or "").replace(" ", "").lower()


def load_style_seasonality() -> tuple[dict[str, str], dict[str, str], dict[str, str]]:
    """(by Code, by franchise base, by layer) -> Style Seasonality.

    By Code and by base only keep values that are unambiguous (one value
    seen); by layer is the majority across that layer's style rows.
    """
    by_code: dict[str, set] = defaultdict(set)
    by_base: dict[str, set] = defaultdict(set)
    by_layer: dict[str, Counter] = defaultdict(Counter)
    for path in (FW27_PATH, SS28_PATH):
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        for r in ws.iter_rows(min_row=3, values_only=True):
            if r[0] is None or not r[25]:  # style rows only
                continue
            by_code[str(r[2])].add(r[25])
            by_base[lib.base_name(strip_special_edition(r[3])[0])].add(r[25])
            by_layer[r[14]][r[25]] += 1
    for path in PIC_PATHS:
        ws = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0]
        for r in ws.iter_rows(min_row=2, values_only=True):
            if r[5] is None or not str(r[5]).isdigit() or not r[22]:
                continue
            by_code[str(r[5])].add(r[22])
    return ({c: next(iter(v)) for c, v in by_code.items() if len(v) == 1},
            {b: next(iter(v)) for b, v in by_base.items() if len(v) == 1},
            {_layer_key(l): c.most_common(1)[0][0] for l, c in by_layer.items() if l})


def load_unit_values() -> dict[str, tuple[float, float]]:
    """Article No. -> (unit landed cost, unit WP), quantity-weighted across
    the Physical stock and Purchase SS27 sheets."""
    wb = openpyxl.load_workbook(STOCK_PATH, read_only=True, data_only=True)
    acc: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0.0])  # qty, cost, wp
    for r in wb["Physical stock 261029"].iter_rows(min_row=2, values_only=True):
        if r[5] in (None, "Total") or not r[9]:
            continue
        a = acc[str(r[5])]
        a[0] += r[9]; a[1] += r[10] or 0; a[2] += r[11] or 0
    for r in wb["Purchase SS27"].iter_rows(min_row=2, values_only=True):
        if r[2] in (None, "Total") or not r[6]:
            continue
        a = acc[str(r[2])]
        a[0] += r[6]; a[1] += r[7] or 0; a[2] += r[8] or 0
    return {k: (v[1] / v[0], v[2] / v[0]) for k, v in acc.items() if v[0] > 0}


def load_sales_2025() -> dict[str, tuple[float, float]]:
    """Article name (whitespace-normalised) -> (units sold, SEK sales), all channels."""
    ws = openpyxl.load_workbook(SALES_PATH, read_only=True, data_only=True).worksheets[0]
    out: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0])
    for r in ws.iter_rows(min_row=3, values_only=True):
        if not r[5] or r[0] == "Total":
            continue
        o = out[_norm(r[5])]
        o[0] += r[9] or 0; o[1] += r[7] or 0
    return {k: (v[0], v[1]) for k, v in out.items()}


def build() -> list[dict]:
    sku_rows = [r for r in build_inventory()[0] if r["In Active Range"] == "N"]
    sea_code, sea_base, sea_layer = load_style_seasonality()
    unit_values = load_unit_values()
    sales = load_sales_2025()

    # Fallback unit values by colour, then style, for SKUs with no own cost/WP.
    fb: dict[str, list[float]] = defaultdict(lambda: [0.0, 0.0, 0])
    for art_no, (c, w) in unit_values.items():
        for key in (art_no[:9], art_no[:6]):
            fb[key][0] += c; fb[key][1] += w; fb[key][2] += 1

    def unit_value(art_no: str) -> tuple[float, float] | None:
        if art_no in unit_values:
            return unit_values[art_no]
        for key in (art_no[:9], art_no[:6]):
            if fb[key][2]:
                return fb[key][0] / fb[key][2], fb[key][1] / fb[key][2]
        return None

    style_tags: dict[str, set] = defaultdict(set)
    style_units: dict[str, int] = defaultdict(int)
    for r in sku_rows:
        if r["Season Article"]:
            style_tags[r["Style Code"]].add(r["Season Article"][0])
        style_units[r["Style Code"]] += r["Available Stock"]

    def seasonality(r: dict) -> tuple[str, str]:
        code = r["Style Code"]
        if code in sea_code:
            return sea_code[code], "1 own range row"
        if r["Franchise"] in sea_base:
            return sea_base[r["Franchise"]], "2 active successor"
        if style_tags[code] == {"S", "F"}:
            return "All year", "3 own S+F colour tags"
        if _layer_key(r["Layer"]) in sea_layer:
            return sea_layer[_layer_key(r["Layer"])], "4 layer default (Provisional)"
        return "Unknown", "none"

    colour_layer: dict[str, str] = {}
    colour_season: dict[str, str] = {}
    for r in sku_rows:
        if r["Layer"]:
            colour_layer.setdefault(r["Article No."][:9], r["Layer"])
        if r["Season Article"]:
            colour_season.setdefault(r["Article No."][:9], r["Season Article"])
    for r in sku_rows:
        r["Layer"] = colour_layer.get(r["Article No."][:9], "")
        r["Season Article"] = colour_season.get(r["Article No."][:9])

    colours: dict[str, dict] = {}
    for r in sku_rows:
        key = r["Article No."][:9]
        c = colours.get(key)
        if c is None:
            sea, src = seasonality(r)
            c = colours[key] = {
                "Style Code": r["Style Code"], "Article": r["Article"], "Color": r["Color"],
                "Franchise": r["Franchise"], "Layer": r["Layer"],
                "Season Article": r["Season Article"], "Age": age_bucket(r["Season Article"]),
                "Style Seasonality": sea, "Seasonality Source": src,
                "Available Stock": 0, "Landed Cost SEK": 0.0, "WP SEK": 0.0,
                "_sizes": 0, "_sizes_in_stock": 0, "_no_value_units": 0,
            }
        c["Available Stock"] += r["Available Stock"]
        if r["Size"] != "1-SIZE":
            c["_sizes"] += 1
            c["_sizes_in_stock"] += r["Available Stock"] > 0
        uv = unit_value(r["Article No."])
        if uv:
            c["Landed Cost SEK"] += r["Available Stock"] * uv[0]
            c["WP SEK"] += r["Available Stock"] * uv[1]
        else:
            c["_no_value_units"] += r["Available Stock"]

    rows = []
    for key, c in colours.items():
        if c["Available Stock"] <= 0:
            continue
        units_25, sek_25 = sales.get(_norm(c["Article"]), (0.0, 0.0))
        weekly = units_25 / 52
        reliable = sek_25 >= lib.MIN_RELIABLE and weekly > 0
        cover = style_units[c["Style Code"]] / weekly if reliable else None
        row = {
            "Style + Colour": key,
            **{k: v for k, v in c.items() if not k.startswith("_")},
            "Markdown Headroom": (1 - c["Landed Cost SEK"] / c["WP SEK"]) if c["WP SEK"] else None,
            "Sizes In Stock / Listed": (f"{c['_sizes_in_stock']}/{c['_sizes']}" if c["_sizes"] else "1-SIZE"),
            "Style 2025 Units Sold": units_25,
            "Style 2025 Sales SEK": round(sek_25),
            "Style Weeks of Cover": round(cover, 1) if cover is not None else None,
        }
        for label, start in WINDOWS.items():
            weeks = (WINDOW_END - start).days / 7
            row[f"Clears by 1-Sep {label}"] = ("" if cover is None
                                               else "Y" if cover <= weeks else "N")
        row["Landed Cost SEK"] = round(row["Landed Cost SEK"])
        row["WP SEK"] = round(row["WP SEK"])
        rows.append(row)
    rows.sort(key=lambda r: (AGE_BUCKETS.index(r["Age"]), -r["WP SEK"]))
    return rows


def matrix(rows: list[dict], row_key: str, row_order: list[str] | None = None) -> list[dict]:
    """row_key x Style Seasonality, units/cost/WP/headroom per cell."""
    seasons = ["All year", "Fall/Winter", "Spring/Summer", "Unknown"]
    cells: dict[tuple, list[float]] = defaultdict(lambda: [0, 0.0, 0.0])
    for r in rows:
        for k in ((r[row_key], r["Style Seasonality"]), (r[row_key], "Total")):
            cells[k][0] += r["Available Stock"]
            cells[k][1] += r["Landed Cost SEK"]
            cells[k][2] += r["WP SEK"]
    keys = row_order or sorted({r[row_key] for r in rows},
                               key=lambda k: -cells[(k, "Total")][2])
    total_units = sum(r["Available Stock"] for r in rows)
    out = []
    for k in keys + ["Total"]:
        line = {row_key: k}
        for s in seasons + ["Total"]:
            if k == "Total":
                u = sum(cells[(kk, s)][0] for kk in keys)
                cost = sum(cells[(kk, s)][1] for kk in keys)
                wp = sum(cells[(kk, s)][2] for kk in keys)
            else:
                u, cost, wp = cells[(k, s)]
            line[f"{s} -- Units"] = u
            if s == "Total":
                line["Share of Units"] = u / total_units if total_units else None
                line["Landed Cost SEK"] = round(cost)
                line["WP SEK"] = round(wp)
                line["Markdown Headroom"] = (1 - cost / wp) if wp else None
        out.append(line)
    return out


def _write_sheet(wb, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    headers = list(dict.fromkeys(h for r in rows for h in r))
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for r in rows:
        ws.append([r.get(h) for h in headers])
    for i, h in enumerate(headers, start=1):
        if h in ("Markdown Headroom", "Share of Units"):
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "0.0%"
        elif "SEK" in h or "Units" in h or h == "Available Stock":
            for cell in ws[get_column_letter(i)][1:]:
                cell.number_format = "#,##0"
        width = max(len(h), max((len(str(r.get(h, ""))) for r in rows), default=0)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions


if __name__ == "__main__":
    rows = build()
    out_path = OUTPUTS_DIR / "SS27_Clearance_View.xlsx"
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Age x Seasonality", matrix(rows, "Age", AGE_BUCKETS))
    _write_sheet(wb, "Layer x Seasonality", matrix(rows, "Layer"))
    _write_sheet(wb, "Colour Detail", rows)
    wb.save(out_path)

    units = sum(r["Available Stock"] for r in rows)
    print(f"{len(rows)} colours, {units} units written to {out_path}")
