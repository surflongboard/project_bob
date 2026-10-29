"""
Builds a generation-by-generation timeline per franchise, for the Nov 12
workshop's lifecycle discussion -- "which products have had a generation
update, and roughly when."

Source: outputs/SS27_Inventory_Review.xlsx's "SKU Detail" sheet (itself
from inputs/stock/Stock_and_purchase_261029.xlsx), grouped by (Franchise
base, Style Code, Article name) -- each distinct group is one generation
of that franchise, whether it's the current Active one or an older one
still carrying leftover stock (RR-019/RR-024).

IMPORTANT LIMITATION, flagged rather than silently presented as fact:
there is no "launch date" or "introduced" field anywhere in this
workstream's sources. The best available proxy is each SKU's own
"Season Article" tag in the stock file (S22 through F27 observed) --
the earliest season seen for a style Code's SKUs is used here as an
approximate "introduced by at least this season" marker, NOT a
confirmed launch date:
  - It only reflects SKUs that still appear in this one stock/purchase
    snapshot (physical stock, undelivered purchase, available stock, or
    orderbook) as of 29-Oct-2026. A generation that sold through
    completely with zero remaining stock anywhere in this file is
    INVISIBLE here, even if it existed -- this timeline is a floor on
    history, not a complete one.
  - 31% of SKU rows have no Season Article at all (blank in source).
  - A season tag may reflect a later re-buy/production run of an
    existing style rather than its true first-ever season.
Treat the "Earliest/Latest Season Seen" columns as a rough signal to
sanity-check in the workshop discussion, not a sourced fact -- per this
register's Open/Provisional/Confirmed discipline, this entire script's
output is Provisional.

Generation ordering within a franchise: no version token < "II"/"2.0"
(2nd gen -- different product lines use one convention or the other,
never both, so they're treated as equivalent rank) < "III" < "IV".

Run: python3 build_generation_timeline.py
Output: ../outputs/Franchise_Generation_Timeline.xlsx
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RR_DIR = Path(__file__).resolve().parents[1]
REPO_ROOT = RR_DIR.parents[1]
sys.path.insert(0, str(REPO_ROOT / "data" / "ss26_portfolio_tiering" / "scripts"))
import ss26_lib as lib  # noqa: E402

OUTPUTS_DIR = RR_DIR / "outputs"
INVENTORY_PATH = OUTPUTS_DIR / "SS27_Inventory_Review.xlsx"

VERSION_RANK = {None: 0, "II": 2, "2.0": 2, "III": 3, "IV": 4}


def season_sort_key(season: str | None) -> tuple[int, int] | None:
    if not season or len(season) < 2:
        return None
    letter, year_part = season[0], season[1:]
    if letter not in ("S", "F") or not year_part.isdigit():
        return None
    return (int(year_part), 0 if letter == "S" else 1)


def season_display(key: tuple[int, int] | None) -> str:
    if key is None:
        return ""
    year, half = key
    return ("S" if half == 0 else "F") + f"{year:02d}"


def load_generation_groups() -> list[dict]:
    wb = openpyxl.load_workbook(INVENTORY_PATH, data_only=True)
    ws = wb["SKU Detail"]
    headers = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(headers)}

    groups: dict[tuple[str, str, str], dict] = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        franchise = r[idx["Franchise"]]
        if not franchise:
            continue
        style_code = r[idx["Style Code"]]
        article = (r[idx["Article"]] or "").strip()
        key = (franchise, style_code, article)
        g = groups.setdefault(key, {
            "franchise": franchise, "style_code": style_code, "article": article,
            "layer": r[idx["Layer"]], "in_active_range": r[idx["In Active Range"]],
            "available_stock": 0, "physical_stock": 0, "n_skus": 0,
            "season_keys": [],
        })
        g["available_stock"] += r[idx["Available Stock"]] or 0
        g["physical_stock"] += r[idx["Physical Stock"]] or 0
        g["n_skus"] += 1
        sk = season_sort_key(r[idx["Season Article"]])
        if sk is not None:
            g["season_keys"].append(sk)

    out = []
    for g in groups.values():
        version = lib.version_token(g["article"])
        out.append({
            "Franchise": g["franchise"],
            "Generation": version or "",
            "_rank": VERSION_RANK.get(version, 99),
            "Style Code": g["style_code"],
            "Article": g["article"],
            "Layer": g["layer"],
            "Active in current range": g["in_active_range"],
            "Earliest Season Seen": season_display(min(g["season_keys"])) if g["season_keys"] else "",
            "Latest Season Seen": season_display(max(g["season_keys"])) if g["season_keys"] else "",
            "# SKUs": g["n_skus"],
            "Available Stock": g["available_stock"],
            "Physical Stock": g["physical_stock"],
        })
    out.sort(key=lambda r: (r["Franchise"], r["_rank"], r["Article"]))
    for r in out:
        del r["_rank"]
    return out


def aggregate_by_generation(code_rows: list[dict]) -> list[dict]:
    """Collapses gender/Asia-Fit variants of the same generation (e.g. Men
    + Women + "AF" codes) into one row per (Franchise, Generation) -- the
    level the lifecycle discussion actually operates at, not per style code."""
    groups: dict[tuple[str, str], dict] = {}
    for r in code_rows:
        key = (r["Franchise"], r["Generation"])
        g = groups.setdefault(key, {
            "Franchise": r["Franchise"], "Generation": r["Generation"],
            "Style Codes": [], "Active in current range": False,
            "earliest": None, "latest": None,
            "Available Stock": 0, "Physical Stock": 0,
        })
        g["Style Codes"].append(r["Style Code"])
        if r["Active in current range"] == "Y":
            g["Active in current range"] = True
        if r["Earliest Season Seen"]:
            k = season_sort_key(r["Earliest Season Seen"])
            if g["earliest"] is None or k < g["earliest"]:
                g["earliest"] = k
        if r["Latest Season Seen"]:
            k = season_sort_key(r["Latest Season Seen"])
            if g["latest"] is None or k > g["latest"]:
                g["latest"] = k
        g["Available Stock"] += r["Available Stock"]
        g["Physical Stock"] += r["Physical Stock"]

    out = []
    for g in groups.values():
        out.append({
            "Franchise": g["Franchise"],
            "Generation": g["Generation"],
            "Style Codes": ", ".join(sorted(set(g["Style Codes"]))),
            "Active in current range": "Y" if g["Active in current range"] else "N",
            "Earliest Season Seen": season_display(g["earliest"]),
            "Latest Season Seen": season_display(g["latest"]),
            "Available Stock": g["Available Stock"],
            "Physical Stock": g["Physical Stock"],
        })
    rank = {None: 0, "II": 2, "2.0": 2, "III": 3, "IV": 4, "": 0}
    out.sort(key=lambda r: (r["Franchise"], rank.get(r["Generation"], 99)))
    return out


def write_workbook(all_rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    def write_sheet(title: str, rows: list[dict]) -> None:
        ws = wb.create_sheet(title)
        headers = list(rows[0].keys())
        ws.append(headers)
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="4472C4")
        for r in rows:
            ws.append(list(r.values()))
        for i, h in enumerate(headers, start=1):
            width = max(len(h), max((len(str(r[h])) for r in rows), default=0)) + 2
            ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

    gen_rows = aggregate_by_generation(all_rows)
    by_franchise_gen: dict[str, list[dict]] = defaultdict(list)
    for r in gen_rows:
        by_franchise_gen[r["Franchise"]].append(r)
    multi_gen_bases = {fr for fr, rows in by_franchise_gen.items() if len(rows) > 1}
    multi_gen_franchise_rows = [r for fr, rows in by_franchise_gen.items() if len(rows) > 1 for r in rows]

    # The blank/no-token generation only gets called "Generation 1" within a franchise that
    # actually HAS multiple generations -- a single-generation franchise has nothing to number
    # against, so it stays blank rather than implying a sequence that doesn't exist.
    for r in gen_rows:
        if r["Generation"] == "" and r["Franchise"] in multi_gen_bases:
            r["Generation"] = "Generation 1"
    for r in all_rows:
        if r["Generation"] == "" and r["Franchise"] in multi_gen_bases:
            r["Generation"] = "Generation 1"

    write_sheet("Generations (by franchise)", gen_rows)
    if multi_gen_franchise_rows:
        write_sheet("Multi-Generation Only", multi_gen_franchise_rows)
    write_sheet("All Style Codes (detail)", all_rows)

    notes = wb.create_sheet("Notes")
    notes.append(["What this is", "'Generations (by franchise)' -- one row per (Franchise, Generation), "
                                   "Men/Women/Asia-Fit codes of the same generation collapsed together. "
                                   "'Multi-Generation Only' is the subset of franchises with 2+ distinct "
                                   "generations -- the actual topic for the lifecycle discussion. "
                                   "'All Style Codes (detail)' keeps every individual style Code separate, "
                                   "for drilling into a specific franchise."])
    notes.append(["No launch-date field exists", "There is no 'introduced'/'launched' field in any source "
                                                   "this workstream has. 'Earliest/Latest Season Seen' is a "
                                                   "PROXY from each SKU's own 'Season Article' tag in the "
                                                   "stock file -- the earliest season a style Code's SKUs "
                                                   "carry, not a confirmed first-ever season."])
    notes.append(["Known limitation 1", "Only SKUs still present in this one stock/purchase snapshot "
                                         "(physical stock, undelivered purchase, available stock, or "
                                         "orderbook) as of 29-Oct-2026 show up at all. A generation that "
                                         "sold through completely with zero remaining stock anywhere in "
                                         "this file is invisible here -- this is a FLOOR on generational "
                                         "history, not a complete one."])
    notes.append(["Known limitation 2", "31% of SKU rows have no Season Article tag at all. A season tag "
                                         "may also reflect a later re-buy/production run of an existing "
                                         "style rather than its true first season."])
    notes.append(["Status", "Provisional -- a reasonable methodology for a first look, not independently "
                             "confirmed. Flag for business confirmation before treating any specific date "
                             "as settled."])
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)
    notes.column_dimensions["A"].width = 20
    notes.column_dimensions["B"].width = 100

    wb.save(path)


if __name__ == "__main__":
    rows = load_generation_groups()
    out_path = OUTPUTS_DIR / "Franchise_Generation_Timeline.xlsx"
    write_workbook(rows, out_path)

    gen_rows = aggregate_by_generation(rows)
    by_franchise: dict[str, int] = defaultdict(int)
    for r in gen_rows:
        by_franchise[r["Franchise"]] += 1
    n_multi = sum(1 for v in by_franchise.values() if v > 1)
    print(f"{len(rows)} style-code rows / {len(gen_rows)} generation rows across {len(by_franchise)} franchises written to {out_path}")
    print(f"{n_multi} franchises show 2+ generations in available stock/purchase data")
