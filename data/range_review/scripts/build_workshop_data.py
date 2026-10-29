"""
Merges the franchise review list, MOQ exceptions, and inventory review
into one JSON dataset for the Nov 12 layer-by-layer workshop page
(the web-based review format, chosen over Excel because embedded images
didn't work well there -- RR-015).

Join key is the BASE name (version-stripped, `_base` on each franchise
row), not the displayed "Franchise" column (version-kept, RR-003) --
build_moq_exceptions.py and build_inventory_review.py both key their own
"Franchise" field on the base name already, so joining on the displayed
name would silently miss every franchise with a version token (e.g.
"Alert GTX Jacket" vs "Alert GTX II Jacket" -- the bug this script
exists to not repeat).

Carries both the FW27+SS28 combined forecast (forecastQty/forecastGP, as
shown on the franchise list) and per-season splits (forecastQtyFW27/
forecastGPFW27, forecastQtySS28/forecastGPSS28) computed in
build_franchise_list.py -- the workshop page's FW27/SS28/All toggle shows
whichever one matches the selected season rather than always showing the
combined total, which it did at first and was misleading once the toggle
existed.

Run: python3 build_workshop_data.py
Output: ../outputs/workshop_data.json
"""
from __future__ import annotations

import json
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl

RR_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RR_DIR / "scripts"))
from build_franchise_list import build_rows  # noqa: E402
from build_forecast_discrepancies import load_ss28_qty_by_code  # noqa: E402

REPO_ROOT = RR_DIR.parents[1]
sys.path.insert(0, str(REPO_ROOT / "data" / "ss26_portfolio_tiering" / "scripts"))
import ss26_lib as lib  # noqa: E402

OUTPUTS_DIR = RR_DIR / "outputs"
MOQ_PATH = OUTPUTS_DIR / "SS28_MOQ_Exceptions.xlsx"
INVENTORY_PATH = OUTPUTS_DIR / "SS27_Inventory_Review.xlsx"


def load_moq_by_base() -> dict[str, list[dict]]:
    """Each MOQ-exception row also carries the SS28 forecast CSV's own Qty
    for the same style Code, and whether it disagrees with the MOQ email's
    FC1 demand (RR-023) -- surfaced as its own "SS28 forecast discrepancy"
    flag/toggle on the board, separate from the MOQ shortfall itself."""
    ss28_csv_qty_by_code = load_ss28_qty_by_code()
    wb = openpyxl.load_workbook(MOQ_PATH, data_only=True)
    ws = wb["SS28 MOQ Exceptions"]
    by_base: dict[str, list[dict]] = defaultdict(list)
    for r in ws.iter_rows(min_row=2, values_only=True):
        base = r[2]
        code, demand_fc1 = r[0], r[11]
        csv_qty = ss28_csv_qty_by_code.get(code)
        has_discrepancy = csv_qty is None or round(csv_qty) != demand_fc1
        by_base[base].append({
            "code": code, "style": r[1], "colorway": r[6],
            "moq_style": r[9], "demand_fc1": demand_fc1, "shortfall": r[12],
            "buying_comment": r[13], "merch_comment": r[14],
            "ss28CsvQty": round(csv_qty) if csv_qty is not None else None,
            "hasDiscrepancy": has_discrepancy,
        })
    return dict(by_base)


def load_inventory_by_base() -> dict[str, dict]:
    wb = openpyxl.load_workbook(INVENTORY_PATH, data_only=True)
    ws = wb["Franchise Summary"]
    by_base = {}
    for r in ws.iter_rows(min_row=2, values_only=True):
        base = r[0]
        by_base[base] = {
            "available_active": r[3],
            "available_not_active": r[4],
            "skus_total": r[5],
            "skus_not_in_range": r[6],
        }
    return by_base


def load_tail_styles_by_base() -> dict[str, list[dict]]:
    """Per franchise, which NOT-in-active-range style code(s) make up its
    inventory tail -- e.g. "Mimic Alert 2.0 Hood" (current) carries leftover
    stock attributed to it because the pre-"2.0" generation ("Mimic Alert
    Hood", codes 607444/607447) shares the same version-stripped base. The
    generic "N units in codes no longer in the active range" alert gave no
    way to see WHICH old style that was -- this breaks it out by style so
    the card can name it, same as the MOQ alert names each flagged style."""
    wb = openpyxl.load_workbook(INVENTORY_PATH, data_only=True)
    ws = wb["SKU Detail"]
    totals: dict[tuple[str, str, str], int] = defaultdict(int)
    for r in ws.iter_rows(min_row=2, values_only=True):
        in_active_range, franchise, style_code, article, available_stock = r[6], r[7], r[1], r[2], r[9]
        if in_active_range == "Y" or not franchise:
            continue
        key = (franchise, style_code, (article or "").strip())
        totals[key] += available_stock or 0

    by_base: dict[str, list[dict]] = defaultdict(list)
    for (base, style_code, article), stock in totals.items():
        if stock <= 0:
            continue
        by_base[base].append({"code": style_code, "style": article, "availableStock": stock})
    for base in by_base:
        by_base[base].sort(key=lambda s: -s["availableStock"])
    return dict(by_base)


def build() -> list[dict]:
    franchise_rows = build_rows()
    moq_by_base = load_moq_by_base()
    inventory_by_base = load_inventory_by_base()
    tail_styles_by_base = load_tail_styles_by_base()

    out = []
    for r in franchise_rows:
        base = r["_base"]
        inv = inventory_by_base.get(base, {})
        moq = moq_by_base.get(base, [])
        franchise_version = lib.version_token(r["Franchise"])
        tail_styles = []
        for s in tail_styles_by_base.get(base, []):
            tail_styles.append({**s, "sameGeneration": lib.version_token(s["style"]) == franchise_version
                                 and franchise_version is not None})
        out.append({
            "base": base,
            "franchise": r["Franchise"],
            "layer": r["Layer"],
            "businessArea": r["Business Area"],
            "gender": r["Gender"],
            "genderNote": r["Gender note"],
            "activeFW27": r["Active in FW27"] == "Y",
            "activeSS28": r["Active in SS28"] == "Y",
            "targetRRP": r["Target RRP"],
            "rrpVaries": r["RRP varies by style"] == "Y",
            "consumerActivity": r["Consumer Activity"],
            "marginTarget": r["Target GM0%"] or None,
            "specialEditionNote": r["Special edition / fit note"],
            "numStyles": r["# styles rolled up"],
            "numColors": r["# active colors"],
            "coreAssortmentFW27": r["Core Assortment FW27"] == "Yes",
            "coreAssortmentSS28": r["Core Assortment SS28"],
            "abs": r["ABS"] == "Yes",
            "asiaFit": r["Asia Fit"] == "Yes",
            "statusFW27": r["Status FW27"],
            "statusSS28": r["Status SS28"],
            "forecastQty": r["Forecast Qty (total)"] or None,
            "forecastGP": r["Forecast GP% (blended)"] or None,
            "forecastQtyDTC": r["Forecast Qty (DTC)"] or None,
            "forecastGPDTC": r["Forecast GP% (DTC)"] or None,
            "forecastQtyWS": r["Forecast Qty (Wholesale)"] or None,
            "forecastGPWS": r["Forecast GP% (Wholesale)"] or None,
            "forecastQtyFW27": r["_forecast_fw27_qty"],
            "forecastGPFW27": r["_forecast_fw27_gp"],
            "forecastQtySS28": r["_forecast_ss28_qty"],
            "forecastGPSS28": r["_forecast_ss28_gp"],
            "imageFile": r["_image_file"],
            "moq": moq,
            "hasForecastDiscrepancy": any(m["hasDiscrepancy"] for m in moq),
            "availableStockActive": inv.get("available_active", 0),
            "availableStockTail": inv.get("available_not_active", 0),
            "skusNotInRange": inv.get("skus_not_in_range", 0),
            "tailStyles": tail_styles,
        })
    return out


if __name__ == "__main__":
    data = build()
    out_path = OUTPUTS_DIR / "workshop_data.json"
    with open(out_path, "w") as f:
        json.dump(data, f, indent=2, ensure_ascii=False)

    n_moq = sum(1 for d in data if d["moq"])
    n_discrepancy = sum(1 for d in data if d["hasForecastDiscrepancy"])
    n_tail = sum(1 for d in data if d["availableStockTail"])
    n_images = sum(1 for d in data if d["imageFile"])
    n_same_gen = sum(1 for d in data for s in d["tailStyles"] if s["sameGeneration"])
    print(f"{len(data)} franchises written to {out_path}")
    print(f"{n_moq} with MOQ flags ({n_discrepancy} with a forecast discrepancy), "
          f"{n_tail} with inventory tail, {n_images} with images")
    print(f"{n_same_gen} tail-style lines are a retired code of the SAME generation, not an older one")
