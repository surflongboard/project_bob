"""
Builds a combined FW27+SS28 style-level forecast summary: volume (Qty),
financials (Gross Sales SEK, COGS SEK, GP%), and a readable Status column
(Carry Over / New / Update), from the two bottom-up forecast files Tobias
supplied separately. See RR-011 (SS28 sourcing) and RR-012 (FW27 sourcing
-- the two files have COMPLETELY different column layouts, not just
different letters).

FW27 source: inputs/forecast/FW27_FC3_data.xlsb.xlsx, sheet "data".
  QTY=col M, GROSS SALES SEK=col O, COGS SEK=col Q, STATUS=col L.
  Column A stacks FOUR forecast vintages (FC1/FC2/FC3/F26) in one sheet --
  only "FC3" (the current forecast, matching the filename) is used here;
  F26 (last comparable season actuals) and FC1/FC2 (earlier forecast
  rounds) are dropped, not summed in by mistake (RR-012).

SS28 source: inputs/forecast/SS28_Forecast_1_Bottom_UpForecast_data.csv.
  QTY=col H, GROSS_SALES_SEK=col Z, COGS_SEK=col AA (RR-011 -- user said
  "ZZ", actual column is AA). 13,796 blank trailing rows filtered out via
  SEASON=='S28', not row position.

Status mapping (RR-013): source STATUS/STATUS_STYLE is CO/NEW/UPDATE in
both files (FW27 FC3 also has a 4th value, NIS, but only in the F26 rows
this script drops). UPDATE is kept as its own label rather than folded
into Carry Over or New -- the user only named two buckets, so a 3rd
distinct value is surfaced rather than silently reclassified.

Shows GP% both blended and split by channel (DTC vs Wholesale) -- the
blended number hides a real spread (SS28 sanity check: DTC 61.3% vs
Wholesale 45.4%), so a reviewer deciding "blend or split" needs both
numbers in front of them, not just the blend. FW27's SALES_CHANNEL
(WS/DTC) and SS28's CHANNEL (WHOLESALE/DTC) are normalized to the same
two labels here.

Run: python3 build_forecast_summary.py
Output: ../outputs/FW27_SS28_Forecast_Summary.xlsx
"""
from __future__ import annotations

import csv
import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

RR_DIR = Path(__file__).resolve().parents[1]
INPUTS_DIR = RR_DIR / "inputs" / "forecast"
OUTPUTS_DIR = RR_DIR / "outputs"
FW27_PATH = INPUTS_DIR / "FW27_FC3_data.xlsb.xlsx"
SS28_PATH = INPUTS_DIR / "SS28_Forecast_1_Bottom_UpForecast_data.csv"

STATUS_LABELS = {"CO": "Carry Over", "NEW": "New", "UPDATE": "Update"}
CHANNEL_LABELS = {"WS": "Wholesale", "WHOLESALE": "Wholesale", "DTC": "DTC"}


def load_fw27(path: Path) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb["data"]
    rows = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] != "FC3":  # drop FC1/FC2 (earlier forecast rounds) and F26 (prior season actuals)
            continue
        rows.append({
            "season": "FW27", "code": r[3], "style": r[4], "layer": r[2],
            "business_area": r[1], "gender": r[7], "status": r[11],
            "channel": CHANNEL_LABELS.get(r[20], r[20]),
            "qty": r[12] or 0, "gross_sales_sek": r[14] or 0, "cogs_sek": r[16] or 0,
        })
    return rows


def _num(s: str) -> float:
    s = s.replace(",", "").strip()
    return float(s) if s else 0.0


def load_ss28(path: Path) -> list[dict]:
    rows = []
    with open(path, newline="", encoding="cp1252") as f:
        for r in csv.DictReader(f):
            if r["SEASON"] != "S28":  # drops the 13,796 blank trailing rows
                continue
            rows.append({
                "season": "SS28", "code": r["CODE"], "style": r["STYLE_NAME"],
                "layer": r["CATEGORY"], "business_area": r["BA"], "gender": r["GENDER"],
                "status": r["STATUS_STYLE"],
                "channel": CHANNEL_LABELS.get(r["CHANNEL"], r["CHANNEL"]),
                "qty": _num(r["QTY"]),
                "gross_sales_sek": _num(r["GROSS_SALES_SEK"]), "cogs_sek": _num(r["COGS_SEK"]),
            })
    return rows


def build_rows() -> list[dict]:
    rows = load_fw27(FW27_PATH) + load_ss28(SS28_PATH)

    by_style: dict[tuple[str, str], list[dict]] = defaultdict(list)
    for r in rows:
        by_style[(r["season"], r["code"])].append(r)

    def totals(items: list[dict]) -> tuple[float, float, float, float | None]:
        qty = sum(r["qty"] for r in items)
        gross = sum(r["gross_sales_sek"] for r in items)
        cogs = sum(r["cogs_sek"] for r in items)
        gp_pct = (gross - cogs) / gross * 100 if gross else None
        return qty, gross, cogs, gp_pct

    out_rows = []
    for (season, code), items in by_style.items():
        first = items[0]
        qty, gross, cogs, gp_pct = totals(items)
        dtc_qty, dtc_gross, dtc_cogs, dtc_gp = totals([r for r in items if r["channel"] == "DTC"])
        ws_qty, ws_gross, ws_cogs, ws_gp = totals([r for r in items if r["channel"] == "Wholesale"])
        statuses = {r["status"] for r in items}
        status_display = "; ".join(sorted(STATUS_LABELS.get(s, s) for s in statuses))

        out_rows.append({
            "Season": season,
            "Code": code,
            "Style Name": first["style"],
            "Layer": first["layer"],
            "Business Area": first["business_area"],
            "Gender": first["gender"],
            "Status": status_display,
            "Qty (total)": round(qty),
            "Gross Sales SEK (total)": round(gross),
            "COGS SEK (total)": round(cogs),
            "GP% (blended)": round(gp_pct, 1) if gp_pct is not None else None,
            "Qty (DTC)": round(dtc_qty),
            "GP% (DTC)": round(dtc_gp, 1) if dtc_gp is not None else None,
            "Qty (Wholesale)": round(ws_qty),
            "GP% (Wholesale)": round(ws_gp, 1) if ws_gp is not None else None,
        })

    out_rows.sort(key=lambda r: (r["Layer"] or "", r["Season"], r["Style Name"] or ""))
    return out_rows


def write_workbook(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forecast Summary"

    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for r in rows:
        ws.append(list(r.values()))

    for i, h in enumerate(headers, start=1):
        width = max(len(h), max((len(str(r[h])) for r in rows), default=0)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 40)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    wb.save(path)


if __name__ == "__main__":
    rows = build_rows()
    out_path = OUTPUTS_DIR / "FW27_SS28_Forecast_Summary.xlsx"
    write_workbook(rows, out_path)
    print(f"{len(rows)} style rows written to {out_path}")
