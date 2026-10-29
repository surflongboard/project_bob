"""
Builds the SS27 inventory view ("List 2"/"List 3" from the original Nov
12 prep email): which SKUs the current stock/purchase/orderbook position
sits in, and whether that SKU's style is still in the active FW27/SS28
range. Calculated at SKU grain (Article No. = 6-digit style Code + 3-char
Color Code + 3-digit Size Code) per the user's explicit direction --
rolling this up to style or franchise grain first would lose real
per-color/size variation -- then rolled up to a franchise-level summary
for readability.

Source: inputs/stock/Stock_and_purchase_261029.xlsx (Inventory Manager,
29-Oct-2026) -- 4 sheets: Physical stock, Purchase SS27 (undelivered),
Available stock (the headline, already-netted figure), Orderbook. Base
SKU population is the UNION of all 4 sheets' Article Nos, not just
Available stock's, so a SKU with e.g. only an open order and no physical
stock yet still shows up.

"Available stock" is shown AS-IS, with the Inventory Manager's own
caveat carried into this output rather than silently adjusted (RR-018):
it already nets wholesale "hard" orders (the Orderbook) out of Physical
stock + Purchase SS27, but NOT SS27 DTC demand or FW27 allocations taken
from current stock -- both still need to come out, so the figure likely
overstates what's truly free for SS27. The user plans to ask for an SS27
forecast separately to refine this later.

"In Active Range" = the SKU's style Code is Active in the current FW27/
SS28 range files (RR-002/RR-006 scope -- excl. ABS/ABS Exclusive). A SKU
NOT in the active range is a List-3/end-of-life candidate: per the user,
these are "articles to build actions for, by end of SS27."

Run: python3 build_inventory_review.py
Output: ../outputs/SS27_Inventory_Review.xlsx (SKU Detail + Franchise Summary)
"""
from __future__ import annotations

import sys
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

RR_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RR_DIR / "scripts"))
from build_franchise_list import FW27_PATH, SS28_PATH, load_style_rows, strip_special_edition  # noqa: E402

REPO_ROOT = RR_DIR.parents[1]
sys.path.insert(0, str(REPO_ROOT / "data" / "ss26_portfolio_tiering" / "scripts"))
import ss26_lib as lib  # noqa: E402

STOCK_PATH = RR_DIR / "inputs" / "stock" / "Stock_and_purchase_261029.xlsx"
OUTPUTS_DIR = RR_DIR / "outputs"


def load_active_range_lookup() -> dict[str, dict]:
    """Style Code -> {franchise base, layer, business area}, Active main-range only."""
    rows = load_style_rows(FW27_PATH, "FW27") + load_style_rows(SS28_PATH, "SS28")
    main_range = [r for r in rows
                  if r["active"] and r["collection_type"] not in ("ABS", "ABS Exclusive")]
    lookup = {}
    for r in main_range:
        code = str(r["code"])
        if code in lookup:
            continue
        stripped, _ = strip_special_edition(r["style"])
        lookup[code] = {
            "base": lib.base_name(stripped),
            "layer": r["layer"],
            "business_area": r["business_area"],
        }
    return lookup


def load_sku_rows() -> dict[str, dict]:
    """Article No. -> merged fields from all 4 sheets."""
    wb = openpyxl.load_workbook(STOCK_PATH, read_only=True, data_only=True)
    skus: dict[str, dict] = defaultdict(lambda: {
        "article": None, "color": None, "size": None, "season_article": None,
        "stock_file_layer": None,
        "available_stock": 0, "physical_stock": 0, "purchase_ss27": 0, "orderbook": 0,
    })

    ws = wb["Physical stock 261029"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[5] == "Total" or r[5] is None:
            continue
        art_no = str(r[5])
        s = skus[art_no]
        s["article"], s["color"], s["size"] = r[6], r[7], r[8]
        s["stock_file_layer"] = r[2]
        s["physical_stock"] += r[9] or 0

    ws = wb["Purchase SS27"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[2] == "Total" or r[2] is None:
            continue
        art_no = str(r[2])
        s = skus[art_no]
        if s["article"] is None:
            s["article"], s["color"], s["size"] = r[3], r[4], r[5]
        s["purchase_ss27"] += r[6] or 0

    ws = wb["Available stock"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[5] == "Total" or r[5] is None:
            continue
        art_no = str(r[5])
        s = skus[art_no]
        if s["article"] is None:
            s["article"], s["color"], s["size"] = r[6], r[7], r[8]
        if s["stock_file_layer"] is None:
            s["stock_file_layer"] = r[2]
        s["season_article"] = r[9]
        s["available_stock"] += r[10] or 0

    ws = wb["Orderbook"]
    for r in ws.iter_rows(min_row=2, values_only=True):
        if r[0] == "Total" or r[0] is None:
            continue
        art_no = str(r[0])
        s = skus[art_no]
        if s["article"] is None:
            s["article"], s["color"], s["size"] = r[1], r[2], r[3]
        s["orderbook"] += r[4] or 0

    return dict(skus)


def build() -> tuple[list[dict], list[dict]]:
    range_lookup = load_active_range_lookup()
    skus = load_sku_rows()

    sku_rows = []
    for art_no, s in skus.items():
        style_code = art_no[:6]
        range_info = range_lookup.get(style_code)
        if range_info:
            franchise = range_info["base"]
            layer = range_info["layer"]
        else:
            # Not Active in the current range -- still parse a franchise name
            # from the stock file's own Article text (same base_name() logic
            # used everywhere else) rather than leaving it unmapped. This is
            # how an old/discontinued generation still shows which franchise
            # it belonged to.
            stripped, _ = strip_special_edition(s["article"] or "")
            franchise = lib.base_name(stripped) if stripped else ""
            layer = s["stock_file_layer"] or ""

        sku_rows.append({
            "Article No.": art_no,
            "Style Code": style_code,
            "Article": (s["article"] or "").strip(),
            "Color": s["color"],
            "Size": s["size"],
            "Season Article": s["season_article"],
            "In Active Range": "Y" if range_info else "N",
            "Franchise": franchise,
            "Layer": layer,
            "Available Stock": s["available_stock"],
            "Physical Stock": s["physical_stock"],
            "Purchase SS27 (Undelivered)": s["purchase_ss27"],
            "Orderbook (Wholesale, Hard Orders)": s["orderbook"],
        })
    sku_rows.sort(key=lambda r: (r["In Active Range"], r["Franchise"], r["Article No."]))

    franchise_summary: dict[str, dict] = defaultdict(lambda: {
        "layer": "", "in_active_range": "N", "skus_total": 0, "skus_not_in_range": 0,
        "available_active": 0, "available_not_active": 0,
    })
    for r in sku_rows:
        key = r["Franchise"] or "(no franchise match)"
        fs = franchise_summary[key]
        fs["layer"] = r["Layer"] or fs["layer"]
        fs["skus_total"] += 1
        if r["In Active Range"] == "Y":
            fs["in_active_range"] = "Y"
            fs["available_active"] += r["Available Stock"]
        else:
            fs["skus_not_in_range"] += 1
            fs["available_not_active"] += r["Available Stock"]

    summary_rows = []
    for franchise, fs in franchise_summary.items():
        summary_rows.append({
            "Franchise": franchise,
            "Layer": fs["layer"],
            "In Active Range": fs["in_active_range"],
            "Available Stock -- Active-Range SKUs": fs["available_active"],
            "Available Stock -- NOT Active-Range SKUs (List 3 candidate)": fs["available_not_active"],
            "# SKUs total": fs["skus_total"],
            "# SKUs not in active range": fs["skus_not_in_range"],
        })
    summary_rows.sort(key=lambda r: (-r["Available Stock -- NOT Active-Range SKUs (List 3 candidate)"],
                                      r["Layer"], r["Franchise"]))

    return sku_rows, summary_rows


def _write_sheet(wb, title: str, rows: list[dict]) -> None:
    ws = wb.create_sheet(title)
    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    for r in rows:
        ws.append(list(r.values()))
    for i, h in enumerate(headers, start=1):
        width = max(len(h), max((len(str(r[h])) for r in rows), default=0)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 50)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions


def write_workbook(sku_rows: list[dict], summary_rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    _write_sheet(wb, "Franchise Summary", summary_rows)
    _write_sheet(wb, "SKU Detail", sku_rows)
    wb.save(path)


if __name__ == "__main__":
    sku_rows, summary_rows = build()
    out_path = OUTPUTS_DIR / "SS27_Inventory_Review.xlsx"
    write_workbook(sku_rows, summary_rows, out_path)

    n_not_in_range = sum(1 for r in sku_rows if r["In Active Range"] == "N")
    units_not_in_range = sum(r["Available Stock"] for r in sku_rows if r["In Active Range"] == "N")
    distinct_codes_not_in_range = len({r["Style Code"] for r in sku_rows if r["In Active Range"] == "N"})
    print(f"{len(sku_rows)} SKUs written to {out_path}")
    print(f"{n_not_in_range} SKUs ({distinct_codes_not_in_range} distinct style codes) "
          f"not in active range, {units_not_in_range} units of Available Stock")
