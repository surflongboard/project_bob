"""
Cross-checks every SS28 FC1 MOQ-exception demand figure (from Buying's
email, RR-016) against the SS28 forecast CSV's own QTY for the same style
Code -- the comparison that surfaced RR-023 (user spotted a discrepancy
on "Mimic Alert 2.0 Hood": card Forecast Qty SS28 1,778 vs MOQ alert's
demand SS28 FC1 820).

Both sources claim to represent an early/FC1-ish SS28 forecast for the
same SKU, but the SS28 forecast CSV has no vintage column to confirm
that (unlike FW27's file, RR-012). Running this check across all 30 MOQ
exception rows -- not just the one the user happened to look at --
found that NONE of the 30 reconcile: 28 differ outright, 2 codes
(Klyva GTX II Jacket Men/Women) don't even appear in the SS28 forecast
CSV at all.

This is the full list to bring to Tobias / the team for the Nov 12
workshop prep, per RR-023 (Open -- needs Tobias, not a code fix).

Run: python3 build_forecast_discrepancies.py
Output: ../outputs/SS28_Forecast_Discrepancies.xlsx
"""
from __future__ import annotations

import csv
import sys
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RR_DIR = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(RR_DIR / "scripts"))
from build_moq_exceptions import MOQ_TABLE, load_code_lookup  # noqa: E402

INPUTS_DIR = RR_DIR / "inputs" / "forecast"
OUTPUTS_DIR = RR_DIR / "outputs"
SS28_FORECAST_PATH = INPUTS_DIR / "SS28_Forecast_1_Bottom_UpForecast_data.csv"


def load_ss28_qty_by_code() -> dict[str, float]:
    """Same source build_forecast_summary.py uses for the SS28 Forecast Qty
    stat shown on the Range Review Board -- summed per style Code so it's
    comparable to the MOQ email's per-SKU FC1 demand figure."""
    by_code: dict[str, float] = {}
    with open(SS28_FORECAST_PATH, newline="", encoding="cp1252") as f:
        for r in csv.DictReader(f):
            if r["SEASON"] != "S28":
                continue
            code = r["CODE"]
            qty = float((r["QTY"] or "0").replace(",", "").strip() or 0)
            by_code[code] = by_code.get(code, 0.0) + qty
    return by_code


def build_rows() -> list[dict]:
    code_lookup = load_code_lookup()
    ss28_qty_by_code = load_ss28_qty_by_code()

    out_rows = []
    for (code, style, cw_code, cw_name, ss28_range, fw27_range, moq_style,
         moq_color, demand_qty, buying_comment, merch_comment) in MOQ_TABLE:
        info = code_lookup.get(code, {})
        csv_qty = ss28_qty_by_code.get(code)
        if csv_qty is None:
            diff = None
            note = "Code not found in SS28 forecast CSV at all"
        else:
            diff = round(csv_qty - demand_qty)
            note = "Matches" if diff == 0 else ("CSV higher" if diff > 0 else "CSV lower")
        out_rows.append({
            "Code": code,
            "Style": style,
            "Franchise": info.get("base", ""),
            "Layer": info.get("layer", ""),
            "Business Area": info.get("business_area", ""),
            "MOQ Email: Demand FC1 (Qty)": demand_qty,
            "SS28 Forecast CSV: Qty (same Code)": round(csv_qty) if csv_qty is not None else None,
            "Difference": diff,
            "Note": note,
        })
    out_rows.sort(key=lambda r: (r["Layer"] or "", r["Franchise"]))
    return out_rows


def write_workbook(rows: list[dict], path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Forecast Discrepancies"

    headers = list(rows[0].keys())
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="4472C4")
    red_fill = PatternFill("solid", fgColor="FBE2E0")
    for r in rows:
        ws.append(list(r.values()))
        if r["Note"] != "Matches":
            for cell in ws[ws.max_row]:
                cell.fill = red_fill

    for i, h in enumerate(headers, start=1):
        width = max(len(h), max((len(str(r[h])) for r in rows), default=0)) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 45)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    notes = wb.create_sheet("Notes")
    notes.append(["What this is", "Every one of the 30 SS28 FC1 MOQ-exception styles (RR-016, "
                                   "Buying's email), cross-checked against the SS28 forecast CSV's "
                                   "own Qty for the same style Code -- the comparison that found "
                                   "RR-023 (user-spotted discrepancy on 'Mimic Alert 2.0 Hood')."])
    notes.append(["Finding", "NONE of the 30 reconcile. 28 of 30 codes differ outright between "
                              "the two sources; the other 2 (Klyva GTX II Jacket Men/Women) don't "
                              "appear in the SS28 forecast CSV at all. Both sources claim to be an "
                              "early/FC1-ish SS28 forecast for the same SKU -- they are not the "
                              "same snapshot."])
    notes.append(["Why it matters", "The Range Review Board's 'Forecast Qty (SS28)' stat is built "
                                     "from the SS28 forecast CSV, NOT the MOQ email. The two numbers "
                                     "shown on a flagged card (card stat vs MOQ alert) are genuinely "
                                     "different sources and can legitimately disagree -- this is not "
                                     "a bug in the page, but it means neither figure should be read "
                                     "as 'the' SS28 forecast without knowing which round it is."])
    notes.append(["Status", "RR-023, Open. Needs Tobias to confirm which round the SS28 forecast "
                             "CSV actually represents, and why it doesn't match the MOQ email's "
                             "FC1 figures even at single-SKU grain -- not something to resolve by "
                             "picking an assumption in code."])
    for row in notes.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)
    notes.column_dimensions["A"].width = 16
    notes.column_dimensions["B"].width = 100

    wb.save(path)


if __name__ == "__main__":
    rows = build_rows()
    out_path = OUTPUTS_DIR / "SS28_Forecast_Discrepancies.xlsx"
    write_workbook(rows, out_path)
    n_mismatch = sum(1 for r in rows if r["Note"] != "Matches")
    print(f"{len(rows)} MOQ-exception codes checked, {n_mismatch} discrepancies found, written to {out_path}")
