"""Rebuild 4 of Tier Bench's 7 qa_data/*.json files from the master workbook:
franchises_chunk_*.json, generation_comparison.json, channel_mix.json,
tier_summary.json (tier_summary's own numbers are still hand-copied below,
not recomputed -- update them if Sheet 1's summary table changes).

Needs art_channel.pkl (a per-article FY25 Wholesale/Retail/E-com sales+GM
lookup) sitting next to this script -- run `python3 build_art_channel.py`
first if it isn't there yet (that script rebuilds it from the raw SS26
exports; not committed itself, same as this script's own output and
tier_bench.html -- a regenerable build intermediate, not source). Without
it, this script raises FileNotFoundError rather than silently producing
an incomplete channel_mix/generation_comparison.

Does NOT cover stock_by_franchise.json, country_breakdown.json, or
regional_tiering.json -- see extract_stock.py, extract_country_breakdown.py,
and extract_regional_tiering.py for those three instead. Re-run
scripts/update_stock.py / update_country_breakdown.py / update_regional_tiering.py
against a fresh workbook first if those sheets themselves need refreshing.
"""
import openpyxl, pickle, json
from pathlib import Path
from collections import defaultdict

base_dir = Path(__file__).resolve().parent.parent
path = base_dir / "Project_Bob_Portfolio_Tiering_08092026.xlsx"

wb = openpyxl.load_workbook(path, data_only=True)

# ---------- A) Franchises (Sheet 2) ----------
ws2 = wb["2. Full Portfolio (1860)"]
franchises = []
for r in range(5, ws2.max_row + 1):
    base = ws2.cell(row=r, column=1).value
    if base is None: continue
    units25 = ws2.cell(row=r, column=21).value
    units_ytd26 = ws2.cell(row=r, column=9).value
    wholesale_share = ws2.cell(row=r, column=19).value
    franchises.append({
        "base": base, "gender": ws2.cell(row=r, column=2).value,
        "layer": ws2.cell(row=r, column=3).value, "tier": ws2.cell(row=r, column=4).value,
        "origTier": ws2.cell(row=r, column=16).value,
        "sales25": round(ws2.cell(row=r, column=6).value or 0),
        "gm25": ws2.cell(row=r, column=7).value,
        "salesYtd26": round(ws2.cell(row=r, column=8).value or 0),
        "gmYtd26": ws2.cell(row=r, column=10).value,
        "pace": ws2.cell(row=r, column=11).value,
        "growth": ws2.cell(row=r, column=12).value,
        "specialAcct": ws2.cell(row=r, column=13).value,
        "clearanceFlag": ws2.cell(row=r, column=14).value,
        "clearanceDetail": ws2.cell(row=r, column=15).value,
        # REG-015/016/017/018/020/021 -- added to Sheet 2 after this
        # script's first version; folded in 9-Sep-2026 so franchises_chunk
        # stays in sync with the published sheet rather than silently
        # trailing it (caught by diffing this script's output against the
        # committed qa_data/*.json before trusting it -- see git history).
        "coreAssortmentFw27": ws2.cell(row=r, column=17).value,
        "fw27Collection": ws2.cell(row=r, column=18).value,
        "wholesaleSharePct": round(wholesale_share, 1) if wholesale_share is not None else None,
        "channelPattern2025": ws2.cell(row=r, column=20).value,
        "units25": round(units25) if units25 is not None else None,
        "unitsYtd26": round(units_ytd26) if units_ytd26 is not None else None,
        "styleCodes": ws2.cell(row=r, column=22).value,
    })
print("franchises:", len(franchises))

# round GM% fields to 1 decimal
for f in franchises:
    for k in ("gm25","gmYtd26","pace","growth","specialAcct"):
        if f[k] is not None:
            f[k] = round(f[k], 1)

# ---------- B) Generation Detail (Sheet 3) -> rebuild comparison for ALL multi-gen franchises ----------
ws3 = wb["3. Generation Detail"]
rows3 = list(ws3.iter_rows(min_row=8, values_only=True))
by_fr = defaultdict(list)
for r in rows3:
    if r[0] is None: continue
    by_fr[(r[0], r[1])].append(r)

def rank(gen):
    if gen.startswith("Original"): return 0
    if "II" in gen or "2.0" in gen: return 1
    if "III" in gen: return 2
    if "IV" in gen: return 3
    return 9

art_channel_path = Path(__file__).resolve().parent / "art_channel.pkl"
if not art_channel_path.exists():
    raise FileNotFoundError(
        f"{art_channel_path} not found -- run `python3 build_art_channel.py` "
        "first (rebuilds this per-article Wholesale/Retail/E-com sales+GM "
        "pickle from the raw SS26 exports; not committed, a regenerable "
        "build intermediate like this script's own output)."
    )
with open(art_channel_path, "rb") as f:
    art_channel = pickle.load(f)

MIN_RELIABLE = 50000
def channel_breakdown(ch_dict):
    out = {}
    tot = sum(v[0] for v in ch_dict.values())
    for ch in ("Wholesale", "Retail", "E-com"):
        v = ch_dict.get(ch)
        if not v:
            out[ch] = {"sales": 0, "gm": None, "share": 0 if tot else None}
            continue
        s, m = v
        out[ch] = {
            "sales": round(s),
            "gm": round(m/s*100, 1) if abs(s) >= MIN_RELIABLE else None,
            "share": round(s/tot*100, 1) if abs(tot) >= MIN_RELIABLE else None,
        }
    return out

generation_comparison = []
channel_mix_rows = []
for key, recs in by_fr.items():
    recs_sorted = sorted(recs, key=lambda r: rank(r[5]))
    for rec in recs_sorted:
        base, gend, layer, tier, origTier, gen, article, models, seasons, nskus, s25, gm25, s26, gm26, exposure = rec
        ch = art_channel.get((base, gend, article), {})
        cb = channel_breakdown(ch)
        channel_mix_rows.append({
            "base": base, "gender": gend, "article": article, "generation": gen,
            "nSkus": nskus, "sales25": round(s25 or 0), "gm25": round(gm25,1) if gm25 is not None else None,
            "wholesaleSales": cb["Wholesale"]["sales"], "wholesaleGm": cb["Wholesale"]["gm"], "wholesaleShare": cb["Wholesale"]["share"],
            "retailSales": cb["Retail"]["sales"], "retailGm": cb["Retail"]["gm"], "retailShare": cb["Retail"]["share"],
            "ecomSales": cb["E-com"]["sales"], "ecomGm": cb["E-com"]["gm"], "ecomShare": cb["E-com"]["share"],
        })

    if len(recs_sorted) < 2:
        continue
    oldest, newest = recs_sorted[0], recs_sorted[-1]
    gm_old, gm_new = oldest[11], newest[11]
    if gm_old is None or gm_new is None:
        gap = None
    else:
        gap = round(gm_old - gm_new, 1)

    old_ch = art_channel.get((key[0], key[1], oldest[6]), {})
    new_ch = art_channel.get((key[0], key[1], newest[6]), {})
    old_cb = channel_breakdown(old_ch)["Wholesale"]
    new_cb = channel_breakdown(new_ch)["Wholesale"]

    trend = None
    if oldest[13] is not None and newest[13] is not None:
        trend = "holds" if newest[13] < oldest[13] else "reversed"

    verdict = None
    if gap is not None:
        if gap <= 0:
            verdict = "new-gen-same-or-better"
        elif old_cb["gm"] is None or new_cb["gm"] is None:
            verdict = "inconclusive (insufficient wholesale-channel data)"
        elif new_cb["gm"] < old_cb["gm"] - 3:
            verdict = "real-gap (worse even within same channel)"
        else:
            verdict = "mix-driven (channel-mix explains most/all of the gap)"

    generation_comparison.append({
        "base": key[0], "gender": key[1], "layer": oldest[2], "tier": oldest[3], "origTier": oldest[4],
        "oldArticle": oldest[6], "newArticle": newest[6],
        "nSkusOld": oldest[9], "nSkusNew": newest[9],
        "salesOld25": round(oldest[10] or 0), "gmOld25": round(gm_old,1) if gm_old is not None else None,
        "salesNew25": round(newest[10] or 0), "gmNew25": round(gm_new,1) if gm_new is not None else None,
        "gapFy25": gap,
        "salesOld26": round(oldest[12] or 0), "gmOld26": round(oldest[13],1) if oldest[13] is not None else None,
        "salesNew26": round(newest[12] or 0), "gmNew26": round(newest[13],1) if newest[13] is not None else None,
        "trend": trend,
        "wholesaleShareOld": old_cb["share"], "wholesaleShareNew": new_cb["share"],
        "wholesaleGmOld": old_cb["gm"], "wholesaleGmNew": new_cb["gm"],
        "verdict": verdict,
    })

print("generation_comparison:", len(generation_comparison))
print("channel_mix_rows:", len(channel_mix_rows))

# ---------- C) Tier summary (Sheet 1) ----------
tier_summary = [
    {"tier":"Hero + Near-Hero","n":34,"sales25":97931546,"gm25":53.2,"salesYtd26":50733948,"gmYtd26":54.3,"pace":51.8},
    {"tier":"Workhorse + Harvest","n":431,"sales25":316679716,"gm25":48.7,"salesYtd26":139975555,"gmYtd26":47.7,"pace":44.2},
    {"tier":"Problem Child","n":150,"sales25":94513733,"gm25":33.5,"salesYtd26":32955867,"gmYtd26":32.9,"pace":34.9},
    {"tier":"New / Test","n":189,"sales25":122822803,"gm25":46.5,"salesYtd26":92223739,"gmYtd26":51.6,"pace":75.1},
    {"tier":"Thin / Immaterial","n":837,"sales25":6822109,"gm25":28.5,"salesYtd26":58964006,"gmYtd26":47.6,"pace":864.3},
    {"tier":"Exited","n":171,"sales25":7179342,"gm25":40.4,"salesYtd26":4975003,"gmYtd26":36.4,"pace":69.3},
    {"tier":"Clearance — Ex China/Zalando","n":48,"sales25":3810858,"gm25":37.8,"salesYtd26":2668099,"gmYtd26":36.1,"pace":70.0},
    {"tier":"TOTAL","n":1860,"sales25":649760108,"gm25":46.3,"salesYtd26":382496216,"gmYtd26":48.0,"pace":58.9},
]

# ---------- save ----------
out_dir = Path(__file__).resolve().parent / "qa_data"
out_dir.mkdir(exist_ok=True)

# chunk franchises into ~400-row docs
CHUNK = 400
chunks = [franchises[i:i+CHUNK] for i in range(0, len(franchises), CHUNK)]
for i, c in enumerate(chunks):
    with open(out_dir / f"franchises_chunk_{i}.json", "w") as f:
        json.dump({"rows": c}, f)
    print(f"chunk {i}: {len(c)} rows, {len(json.dumps({'rows':c}))} bytes")

with open(out_dir / "generation_comparison.json", "w") as f:
    json.dump({"rows": generation_comparison}, f)
with open(out_dir / "channel_mix.json", "w") as f:
    json.dump({"rows": channel_mix_rows}, f)
with open(out_dir / "tier_summary.json", "w") as f:
    json.dump({"rows": tier_summary}, f)

print("gen_comp bytes:", len(json.dumps({"rows": generation_comparison})))
print("channel_mix bytes:", len(json.dumps({"rows": channel_mix_rows})))
print("n_chunks:", len(chunks))
