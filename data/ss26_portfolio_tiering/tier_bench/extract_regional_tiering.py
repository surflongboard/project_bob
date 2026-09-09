"""
Rebuilds qa_data/regional_tiering.json (Tier Bench's `queryRegionalTiering`
tool) by reading the master workbook's "6. Full Portfolio (Nordic)" and
"7. Full Portfolio (Non-Nordic)" sheets directly -- NOT by re-deriving the
region-relative tiering rule (that logic lives in
scripts/update_regional_tiering.py, REG-023/024, and stays the single
source of truth for Sheets 6/7 themselves). This script's job is only the
last step: turn the published sheets into Tier Bench's JSON shape, so
re-run scripts/update_regional_tiering.py FIRST if the regional sheets
themselves need refreshing, then this script.

One row per franchise per region (3,720 = 1,860 x 2 as of the
08-Sep-2026 workbook) -- Sheets 6/7 are denormalized to one row per
franchise PER COUNTRY (or a single "no regional sales" placeholder row),
so this dedupes by taking the first row per (base, gender) within each
region sheet; the franchise-level columns (Regional Tier, Sales_2024/25,
Units_2025, GM%, Growth%, Wholesale/DTC) are identical across every
country row for that franchise, only the trailing country columns vary
-- this script drops those trailing country columns entirely (Tier
Bench's queryByCountry, extract_country_breakdown.py, already covers
that grain from Sheet 5). `globalTier` is joined in from Sheet 2's own
Tier column, for direct regional-vs-global comparison.

Usage (from this directory):
    python3 extract_regional_tiering.py
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import ss26_lib as lib

REGION_SHEETS = [
    ("6. Full Portfolio (Nordic)", "Nordic"),
    ("7. Full Portfolio (Non-Nordic)", "Non-Nordic"),
]


def main():
    wb = openpyxl.load_workbook(lib.WORKBOOK_PATH, read_only=True, data_only=True)

    ws2 = wb["2. Full Portfolio (1860)"]
    global_tier = {}
    for row in ws2.iter_rows(min_row=lib.FIRST_DATA_ROW, values_only=True):
        base = row[0]
        if base is None:
            continue
        global_tier[(base, row[1])] = row[3]

    out_rows = []
    for sheet_name, region in REGION_SHEETS:
        ws = wb[sheet_name]
        seen = set()
        for row in ws.iter_rows(min_row=lib.FIRST_DATA_ROW, values_only=True):
            base = row[0]
            if base is None:
                continue
            key = (base, row[1])
            if key in seen:
                continue
            seen.add(key)
            sales24, sales25, units25, gm25, growth, wholesale25, dtc25 = row[4:11]
            out_rows.append({
                "base": base,
                "gender": row[1],
                "layer": row[2],
                "region": region,
                "regionalTier": row[3],
                "globalTier": global_tier.get(key),
                "sales24": round(sales24) if sales24 is not None else None,
                "sales25": round(sales25) if sales25 is not None else None,
                "units25": round(units25) if units25 is not None else None,
                "gm25": round(gm25, 1) if gm25 is not None else None,
                "growth": round(growth, 1) if growth is not None else None,
                "wholesale25": round(wholesale25) if wholesale25 is not None else None,
                "dtc25": round(dtc25) if dtc25 is not None else None,
            })
        print(f"{region}: {len(seen)} franchises")
    wb.close()

    out_dir = Path(__file__).resolve().parent / "qa_data"
    out_dir.mkdir(exist_ok=True)
    with open(out_dir / "regional_tiering.json", "w") as f:
        json.dump({"rows": out_rows}, f)
    print("written", out_dir / "regional_tiering.json", len(json.dumps({"rows": out_rows})), "bytes")


if __name__ == "__main__":
    main()
