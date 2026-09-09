"""
Rebuilds qa_data/stock_by_franchise.json (Tier Bench's `queryStock` tool)
from the raw warehouse stock snapshot, using the EXACT same matching logic
as scripts/update_stock.py (REG-019) -- same source file, same column
indices, same franchise-key parsing (`ss26_lib.franchise_key`) -- so this
stays consistent with Sheet 4 rather than drifting from it.

Franchise-level, one row per franchise with any matched stock (656 of
1,860 as of the 08-Sep-2026 workbook) -- NOT the same as Sheet 4's own
detail tables, which cap at the top 10 per tier (REG-019's rerunnable
version deliberately caps there; this script needs every matched
franchise since it's Tier Bench's underlying dataset, not a printed
table). `deadStock` uses the same DEAD_STOCK_SALES_THRESHOLD (1,000 SEK)
as update_stock.py's own dead-stock signal.

Usage (from this directory):
    python3 extract_stock.py [--source PATH]
"""
from __future__ import annotations

import argparse
import json
import sys
from collections import defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import ss26_lib as lib

ARTICLE_COL_IDX = 6   # 0-indexed -- see update_stock.py's docstring ASSUMPTION
STOCK_COL_IDX = 10    # 0-indexed
DATA_MIN_ROW = 3
DEAD_STOCK_SALES_THRESHOLD = 1000  # SEK -- matches update_stock.py exactly

TIER_ORDER = lib.TIER_ORDER


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(lib.INPUTS_DIR / "stock" / "Available_stock_260825.xlsx"))
    ap.add_argument("--sheet", default=None)
    args = ap.parse_args()

    import openpyxl

    wb_stock = openpyxl.load_workbook(args.source, data_only=True)
    if args.sheet:
        ws_stock = wb_stock[args.sheet]
    else:
        candidates = [n for n in wb_stock.sheetnames if n.lower().startswith("available stock")]
        ws_stock = wb_stock[candidates[0]] if candidates else wb_stock[wb_stock.sheetnames[0]]
    rows = list(ws_stock.iter_rows(min_row=DATA_MIN_ROW, values_only=True))

    wb, ws2 = lib.load_full_portfolio()
    pub_lookup = {}
    for r in range(lib.FIRST_DATA_ROW, ws2.max_row + 1):
        b = ws2.cell(row=r, column=1).value
        if b is None:
            continue
        g = (ws2.cell(row=r, column=2).value or "").strip()
        pub_lookup[(b.strip(), g)] = {
            "tier": ws2.cell(row=r, column=4).value,
            "sales25": ws2.cell(row=r, column=6).value or 0,
        }

    franchise_units = defaultdict(float)
    total_rows = matched_rows = 0
    for row in rows:
        article = row[ARTICLE_COL_IDX]
        stock = row[STOCK_COL_IDX] or 0
        if not article:
            continue
        total_rows += 1
        key = lib.franchise_key(article.strip())
        if key not in pub_lookup:
            continue
        franchise_units[key] += stock
        matched_rows += 1

    total_units = sum(row[STOCK_COL_IDX] or 0 for row in rows if row[ARTICLE_COL_IDX])
    matched_units = sum(franchise_units.values())
    print(f"stock rows: {total_rows}  matched an existing franchise: {matched_rows}")
    print(f"total units: {total_units:,.0f}  matched units: {matched_units:,.0f} "
          f"({matched_units / total_units * 100 if total_units else 0:.1f}%)")

    out_rows = []
    for (base, gender), units in franchise_units.items():
        info = pub_lookup[(base, gender)]
        sales25 = round(info["sales25"])
        out_rows.append({
            "base": base,
            "gender": gender,
            "tier": info["tier"],
            "unitsInStock": round(units),
            "sales25": sales25,
            "deadStock": sales25 < DEAD_STOCK_SALES_THRESHOLD,
        })
    out_rows.sort(key=lambda r: -r["unitsInStock"])
    print("franchises with matched stock:", len(out_rows))

    out_dir = Path(__file__).resolve().parent / "qa_data"
    out_dir.mkdir(exist_ok=True)
    with open(out_dir / "stock_by_franchise.json", "w") as f:
        json.dump({"rows": out_rows}, f)
    print("written", out_dir / "stock_by_franchise.json", len(json.dumps({"rows": out_rows})), "bytes")


if __name__ == "__main__":
    main()
