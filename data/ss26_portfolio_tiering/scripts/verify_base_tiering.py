"""
REG-024: validates Sheet 2's published Tier column against the now-confirmed
base 8-tier commercial portfolio rule (business-supplied 9-Sep-2026 --
previously undocumented anywhere in this repo; an earlier attempt to
reverse-engineer it from Sales/GM%/Growth/Pace alone had failed).

This is a READ-ONLY audit -- it does NOT rewrite Sheet 2's Tier column.
Sheet 2's published tiers stay authoritative even though this script can
now reproduce them, because:
  (a) the tiny remaining gap (4 of 1,812 non-Clearance franchises) traces to
      article-name parsing edge cases in THIS reconstruction, not necessarily
      an error in the original -- overwriting published, audited tiers with
      a 99.8%-but-not-100%-matching recomputation would be a regression, not
      an improvement.
  (b) the original build may have had access to source data more precise
      than the bob_salesdata_2024/2025.xlsx reconstruction used here (which
      carries REG-009's own known ~2.3% cross-workstream gap vs. the SS26
      exports).

Run this after any future data refresh to re-confirm the rule still holds
before trusting it for something new (e.g. REG-023's regional tiering,
which DOES use this rule to build new sheets from scratch, since those
don't have a pre-existing "audited" value to protect).

Usage:
    python3 data/ss26_portfolio_tiering/scripts/verify_base_tiering.py
"""
from __future__ import annotations

from collections import defaultdict

import ss26_lib as lib

# Global, non-region-scaled floors -- the exact numbers from the business's
# rule (see 1.2 "The full 8-tier commercial portfolio").
HERO_SALES_FLOOR = 2_000_000
NEAR_HERO_SALES_FLOOR = 1_000_000
WHOLESALE_FLOOR = 100_000
DTC_FLOOR = 50_000

BOB_SALES_FILES = {
    2024: lib.REPO_ROOT / "data" / "bob_salesdata_2024.xlsx",
    2025: lib.REPO_ROOT / "data" / "bob_salesdata_2025.xlsx",
}
NON_MARKET_VALUES = {"Total", "Restored"}


def load_franchise_totals(year):
    """(base, gender) -> {sales, units, margin, wholesale, dtc}, Core-scope only."""
    import openpyxl

    path = BOB_SALES_FILES[year]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = ws.iter_rows(values_only=True)
    header = next(rows)
    idx = {n: i for i, n in enumerate(header)}
    agg = defaultdict(lambda: {"sales": 0.0, "units": 0.0, "margin": 0.0, "wholesale": 0.0, "dtc": 0.0})
    for row in rows:
        market = row[idx["Sales Market"]]
        if market in NON_MARKET_VALUES or market is None:
            continue
        if not lib.is_core_customer_group(row[idx["Customer Group"]]):
            continue
        a = row[idx["Article"]]
        if not a:
            continue
        key = lib.franchise_key(a.strip())
        sales = row[idx["Garp SEK Sales"]] or 0
        agg[key]["sales"] += sales
        agg[key]["units"] += row[idx["Units Sold"]] or 0
        agg[key]["margin"] += row[idx["Garp SEK Margin"]] or 0
        # REG-018's precedent: Marketplace grouped into Wholesale, not DTC.
        if row[idx["Sales Channel"]] in ("Wholesale", "Marketplace"):
            agg[key]["wholesale"] += sales
        else:
            agg[key]["dtc"] += sales
    return agg


def predict(fy24, fy25, key):
    a24 = fy24.get(key, {"sales": 0, "units": 0, "margin": 0, "wholesale": 0, "dtc": 0})
    a25 = fy25.get(key, {"sales": 0, "units": 0, "margin": 0, "wholesale": 0, "dtc": 0})
    sales25 = a25["sales"]
    gm25 = (a25["margin"] / sales25 * 100) if sales25 else 0.0
    growth = (sales25 - a24["sales"]) / a24["sales"] * 100 if a24["sales"] else None
    return lib.classify_tier(
        sales=sales25, units_prior=a24["units"], units_current=a25["units"],
        growth=growth, gm_pct=gm25, wholesale_sek=a25["wholesale"], dtc_sek=a25["dtc"],
        sales_floor_hero=HERO_SALES_FLOOR, sales_floor_near_hero=NEAR_HERO_SALES_FLOOR,
        wholesale_floor=WHOLESALE_FLOOR, dtc_floor=DTC_FLOOR,
    )


def main():
    fy24 = load_franchise_totals(2024)
    fy25 = load_franchise_totals(2025)

    wb, ws = lib.load_full_portfolio()
    match = mismatch = skipped_clearance = 0
    mismatches = []
    for r in range(lib.FIRST_DATA_ROW, ws.max_row + 1):
        b = ws.cell(row=r, column=1).value
        if b is None:
            continue
        g = (ws.cell(row=r, column=2).value or "").strip()
        orig_tier = ws.cell(row=r, column=16).value
        clearance_flag = ws.cell(row=r, column=14).value
        if clearance_flag == "Fully":
            skipped_clearance += 1
            continue
        pred = predict(fy24, fy25, (b.strip(), g))
        if pred == orig_tier:
            match += 1
        else:
            mismatch += 1
            mismatches.append((b.strip(), g, orig_tier, pred))

    total = match + mismatch
    print(f"match: {match}/{total} ({match/total*100:.1f}%)  skipped (Clearance, not rule-derived): {skipped_clearance}")
    if mismatches:
        print("mismatches (actual vs predicted):")
        for base, gender, actual_t, pred_t in mismatches:
            print(f"  {base!r} {gender}: actual={actual_t!r}  predicted={pred_t!r}")


if __name__ == "__main__":
    main()
