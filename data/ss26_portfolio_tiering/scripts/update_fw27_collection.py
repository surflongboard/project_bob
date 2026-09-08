"""
REG-016: "FW27 Collection" column (Active / Not active / Not in review)
on the Full Portfolio sheet.

Reads data/Assortment Attribution Review(F27).xlsx (shared with the other
workstream -- Key-Article Margin Analysis -- so it lives at the top-level
data/ folder, not under ss26_portfolio_tiering/inputs/). Column A = Active
(bool), Column C = Style name, data starting row 3.

Classification per franchise (Base, Gender):
  - "Active"        -- the review lists at least one style for this
                        franchise with Active = TRUE.
  - "Not active"     -- the review lists the franchise, but every style
                        under it has Active = FALSE.
  - "Not in review"  -- the franchise doesn't appear in the review at all
                        (user-confirmed label, chosen over blank/"Unknown").

Prints "exceptions" -- franchises marked "Not active" that still carry
material YTD2026 sales (>= MIN_RELIABLE) -- worth flagging to Product,
since REG-016's revision note recorded this exact pattern for franchises
like Bield Down II / Zircon Slim II / Gabbro II.

Usage:
    python3 data/ss26_portfolio_tiering/scripts/update_fw27_collection.py \
        [--source PATH]
"""
from __future__ import annotations

import argparse
from collections import defaultdict

import openpyxl

import ss26_lib as lib


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--source", default=str(lib.REPO_ROOT / "data" / "Assortment Attribution Review(F27).xlsx"))
    args = ap.parse_args()

    wb_ar = openpyxl.load_workbook(args.source, data_only=True)
    ws_ar = wb_ar[wb_ar.sheetnames[0]]
    rows_ar = list(ws_ar.iter_rows(min_row=3, values_only=True))

    franchise_active = defaultdict(set)
    for row in rows_ar:
        active, style = row[0], row[2]
        if not style:
            continue
        key = lib.franchise_key(style)
        franchise_active[key].add(active)

    wb, ws = lib.load_full_portfolio()
    s = lib.styles()

    ws["R4"] = "FW27 Collection"
    ws["R4"].font = s["header_font"]
    ws["R4"].fill = s["header_fill"]
    ws.column_dimensions["R"].width = 16

    counts = {"Active": 0, "Not active": 0, "Not in review": 0}
    exceptions = []
    for r in range(lib.FIRST_DATA_ROW, ws.max_row + 1):
        b = ws.cell(row=r, column=1).value
        if b is None:
            continue
        g = (ws.cell(row=r, column=2).value or "").strip()
        key = (b.strip(), g)
        active_vals = franchise_active.get(key)
        if active_vals is None:
            status = "Not in review"
        elif True in active_vals:
            status = "Active"
        else:
            status = "Not active"
        counts[status] += 1
        ws.cell(row=r, column=18, value=status).font = s["body_font"]

        sales_ytd26 = ws.cell(row=r, column=8).value or 0
        if status == "Not active" and sales_ytd26 >= lib.MIN_RELIABLE:
            exceptions.append((b, g, ws.cell(row=r, column=4).value, sales_ytd26, ws.cell(row=r, column=6).value or 0))

    print("counts:", counts)
    print("exceptions (Not active but material YTD2026 sales):", len(exceptions))
    for e in sorted(exceptions, key=lambda x: -x[3]):
        print(f"  {e[0]:<28} {e[1]:<7} {e[2]:<24} YTD26={e[3]:,.0f}  FY25={e[4]:,.0f}")

    wb.save(lib.WORKBOOK_PATH)
    print("saved")


if __name__ == "__main__":
    main()
