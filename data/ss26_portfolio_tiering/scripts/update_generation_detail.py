"""
REG-014: rebuilds Sheet 3 "Generation Detail" -- one row per distinct
Article, for every franchise that has a genuine NAMING-based version
succession (an "Original" article and a later II/III/IV/2.0 article
under the same Base+Gender).

IMPORTANT (see REG-014's revision note): generation is determined purely
by the article name's version token. It is NOT bucketed by clearance-list
membership -- an earlier build of this sheet conflated the two, which
caused the same article to appear inconsistently depending on which
colorway SKUs happened to be on a close-out list. Clearance Exposure is
shown as a separate, purely informational column; it does not define
which article is "old" vs "new".

Usage:
    python3 data/ss26_portfolio_tiering/scripts/update_generation_detail.py \
        [--jv PATH] [--zalando PATH]

Re-run whenever the SS26 raw exports or the clearance lists change. This
script only rebuilds Sheet 3 -- run update_clearance.py first if the
clearance lists changed, so Sheet 2's Tier column (used here for sort
order) reflects the latest Clearance tier reassignment.
"""
from __future__ import annotations

import argparse
from collections import defaultdict

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

import ss26_lib as lib


def load_clearance(path, sheet=None):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb[wb.sheetnames[0]]
    return list(ws.iter_rows(values_only=True))[1:]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--jv", default=str(lib.INPUTS_DIR / "clearance" / "Close_out_delivery_JV_August_2026.xlsx"))
    ap.add_argument("--zalando", default=str(lib.INPUTS_DIR / "clearance" / "Zalando_close_out_order_August_2026.xlsx"))
    ap.add_argument("--zalando-sheet", default="Zalando close out order August")
    args = ap.parse_args()

    sco_info = {}
    article_scos = defaultdict(set)        # (base, gender, article) -> {SCO, ...}
    franchise_articles = defaultdict(set)  # (base, gender) -> {article, ...}

    def ingest(header, rows, market_field, sales_bucket):
        idx = {n: i for i, n in enumerate(header)}
        for row in rows:
            sco = row[idx["SCO"]]
            article = row[idx["Article"]]
            if sco is None or not article:
                continue
            a = article.strip()
            key = lib.franchise_key(a)
            article_scos[(key[0], key[1], a)].add(sco)
            franchise_articles[key].add(a)
            rec = sco_info.setdefault(sco, {
                "fy25_sales": 0.0, "fy25_margin": 0.0, "ytd26_sales": 0.0, "ytd26_margin": 0.0,
                "article": a, "layer": row[idx["Layer"]], "model": row[idx["Model"]],
                "start_season": row[idx["Start Season"]],
            })
            if lib.is_core_market(row[idx[market_field]]):
                sales = row[idx["Garp SEK Sales"]] or 0
                margin = row[idx["Garp SEK Margin"]] or 0
                rec[f"{sales_bucket}_sales"] += sales
                rec[f"{sales_bucket}_margin"] += margin

    for path in lib.SS26_EXPORTS:
        header, rows = lib.load_raw_export(path)
        market_field = lib.SS26_MARKET_FIELD[path.name]
        bucket = "ytd26" if "SS26YTD" in path.name else "fy25"
        ingest(header, rows, market_field, bucket)

    # clearance SCOs -- informational only, does not define generation
    jv_rows = load_clearance(args.jv)
    za_rows = load_clearance(args.zalando, args.zalando_sheet)
    clearance_scos = defaultdict(set)
    for rows, source in [(jv_rows, "JV/China"), (za_rows, "Zalando")]:
        for row in rows:
            art_no = row[0]
            if art_no is None or art_no == "Total":
                continue
            clearance_scos[str(art_no)[:9]].add(source)

    # franchises with a genuine naming-based version succession
    multi_version_keys = [
        key for key, arts in franchise_articles.items()
        if len({lib.version_token(a) for a in arts}) > 1
    ]
    print("naming-based multi-generation franchises:", len(multi_version_keys))

    article_rows = []
    for key in multi_version_keys:
        base, gend = key
        for art in franchise_articles[key]:
            scos = article_scos[(base, gend, art)]
            vt = lib.version_token(art)
            gen_label = "Original (no version token)" if vt is None else f"Version {vt}"
            s25 = sum(sco_info[s]["fy25_sales"] for s in scos)
            m25 = sum(sco_info[s]["fy25_margin"] for s in scos)
            s26 = sum(sco_info[s]["ytd26_sales"] for s in scos)
            m26 = sum(sco_info[s]["ytd26_margin"] for s in scos)
            gm25 = m25 / s25 * 100 if s25 and abs(s25) >= lib.MIN_RELIABLE else None
            gm26 = m26 / s26 * 100 if s26 and abs(s26) >= lib.MIN_RELIABLE else None
            models = sorted(set(str(sco_info[s]["model"]) for s in scos))
            seasons = sorted(set(
                str(sco_info[s]["start_season"]) for s in scos
                if sco_info[s]["start_season"] and str(sco_info[s]["start_season"]) != "0"
            ))
            flagged, srcs = set(), set()
            for s in scos:
                if s in clearance_scos:
                    flagged.add(s)
                    srcs |= clearance_scos[s]
            layer = sco_info[next(iter(scos))]["layer"]
            article_rows.append(dict(
                key=key, article=art, layer=layer, version_rank=lib.RANK.get(vt, 9), generation=gen_label,
                n_skus=len(scos), sales25=s25, gm25=gm25, sales26=s26, gm26=gm26,
                models=models, seasons=seasons, n_flagged=len(flagged), sources=sorted(srcs),
            ))
    print("total article rows:", len(article_rows))

    # --- write Sheet 3 ---
    wb, ws2 = lib.load_full_portfolio()
    lookup = lib.full_portfolio_lookup(ws2)

    if lib.GENERATION_SHEET in wb.sheetnames:
        del wb[lib.GENERATION_SHEET]
    ws3 = wb.create_sheet(lib.GENERATION_SHEET)

    from openpyxl.styles import Font
    s = lib.styles()
    TITLE_FONT = Font(name="Arial", bold=True, color="FF1F3864")
    WRAP = Alignment(vertical="top", wrap_text=True)

    ws3["A1"] = "Generation Detail — Article Version Succession, by Franchise"
    ws3["A1"].font = TITLE_FONT
    ws3.merge_cells("A1:N1")

    ws3["A3"] = (
        "Generation is determined purely by the article name's version token (e.g. \"X\" = Original vs. \"X II\" = "
        "Version II) — NOT by whether a SKU is on the close-out list. Covers franchises where the raw exports show "
        "a genuine naming-based version succession (an \"Original\" and a later II/III/IV/2.0 article under the "
        "same franchise) — franchises with only one continuously-named article aren't included, since there's no "
        "generation to compare. \"Clearance Exposure\" is shown separately, purely informational — some SKUs of "
        "either generation may or may not also be on a close-out list (REG-010); that overlap does not define "
        "which article is \"old\" or \"new.\""
    )
    ws3["A3"].font = s["body_font"]
    ws3["A3"].alignment = WRAP
    ws3.merge_cells("A3:N3")
    ws3.row_dimensions[3].height = 62

    ws3["A5"] = (
        "CAVEAT (REG-012): Sales_2025/Sales_YTD2026 here are recomputed directly from the raw SS26 exports at "
        "Article/SKU level (same Core-scope rule as Sheet 2) — they are supplementary detail, not the published, "
        "audited Sheet-2 figures, and a franchise's article-rows won't necessarily sum to its Sheet-2 Sales_2025 "
        "(an unresolved raw-export ordertype ambiguity, see REG-012). Read as relative/directional — which "
        "generation is bigger, which has better margin. GM% is blanked below 50,000 SEK of sales."
    )
    ws3["A5"].font = s["gray_font"]
    ws3["A5"].alignment = WRAP
    ws3.merge_cells("A5:N5")
    ws3.row_dimensions[5].height = 62

    headers = ["Base", "Gender", "Layer", "Tier (current)", "Original Tier (pre-consolidation)",
               "Generation", "Article", "Model(s)", "Start Season(s)", "# SKUs",
               "Sales_2025 (SEK, raw)", "GM%_2025 (raw)", "Sales_YTD2026 (SEK, raw)", "GM%_YTD2026 (raw)",
               "Clearance Exposure (informational, REG-010)"]
    hdr_row = 7
    for c, h in enumerate(headers, start=1):
        cell = ws3.cell(row=hdr_row, column=c, value=h)
        cell.font = s["header_font"]
        cell.fill = s["header_fill"]

    tier_rank = {t: i for i, t in enumerate(lib.TIER_ORDER)}

    def sort_key(row):
        info = lookup.get(row["key"], {})
        return (tier_rank.get(info.get("tier"), 99), row["key"][0], row["key"][1], row["version_rank"])

    rows_sorted = sorted(article_rows, key=sort_key)

    r = hdr_row + 1
    for row in rows_sorted:
        base, gend = row["key"]
        info = lookup.get(row["key"], {})
        if row["n_flagged"]:
            exposure = f'{row["n_flagged"]} of {row["n_skus"]} SKUs on close-out list ({" + ".join(row["sources"])})'
        else:
            exposure = "None"
        vals = [
            base, gend, row["layer"], info.get("tier"), info.get("orig_tier"),
            row["generation"], row["article"], "; ".join(row["models"]),
            "/".join(row["seasons"]) if row["seasons"] else None, row["n_skus"],
            row["sales25"], row["gm25"], row["sales26"], row["gm26"], exposure,
        ]
        for c, v in enumerate(vals, start=1):
            cell = ws3.cell(row=r, column=c, value=v)
            cell.font = s["body_font"]
            if c in (11, 13):
                cell.number_format = s["num_fmt"]
            elif c in (12, 14):
                cell.number_format = s["pct_fmt"]
        r += 1

    widths = {1: 24, 2: 9, 3: 12, 4: 20, 5: 22, 6: 24, 7: 32, 8: 20, 9: 14, 10: 8, 11: 16, 12: 12, 13: 16, 14: 12, 15: 42}
    for c, w in widths.items():
        ws3.column_dimensions[get_column_letter(c)].width = w
    ws3.freeze_panes = "A8"

    wb.save(lib.WORKBOOK_PATH)
    print("rows written:", r - hdr_row - 1)


if __name__ == "__main__":
    main()
